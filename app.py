from pathlib import Path
import os, json, subprocess, tempfile, secrets, shutil, sqlite3, re, unicodedata, hashlib
from datetime import datetime, date, timedelta
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, session, abort, send_file
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from io import BytesIO
from PIL import Image as PILImage, ImageOps, ImageDraw
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY','cambiar-esta-clave-en-produccion')
# Color estable por trabajador: respeta el color configurado y evita gris uniforme
# en usuarios históricos que todavía no tienen color guardado.
COLORES_PLANIFICACION = ('#2563eb', '#16a34a', '#d97706', '#9333ea', '#dc2626', '#0891b2', '#be185d', '#4d7c0f', '#b45309', '#0f766e', '#7c3aed', '#1d4ed8', '#c2410c', '#047857', '#a21caf')

@app.template_filter('color_trabajador')
def color_trabajador(usuario):
    if usuario is None:
        return '#64748b'
    valor = (getattr(usuario, 'color', None) or '').strip()
    if len(valor) == 7 and valor.startswith('#') and all(c in '0123456789abcdefABCDEF' for c in valor[1:]):
        return valor
    identificador = getattr(usuario, 'id', None)
    if isinstance(identificador, int) and identificador > 0:
        return COLORES_PLANIFICACION[(identificador - 1) % len(COLORES_PLANIFICACION)]
    return '#64748b'

db_url = os.environ.get('DATABASE_URL', 'sqlite:///ingepro.db')
if db_url.startswith('postgres://'): db_url = db_url.replace('postgres://','postgresql://',1)
app.config['SQLALCHEMY_DATABASE_URI'] = db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)

PERMISOS = {
 'ingreso':'Ingreso de equipos', 'diagnostico':'Diagnóstico', 'cotizaciones':'Cotizaciones',
 'reparacion':'Reparación', 'calidad':'Control de calidad', 'despacho':'Embalaje y despacho',
 'informes':'Informes', 'maestros':'Maestros', 'clientes':'Clientes y contactos', 'planificacion':'Planificación', 'produccion':'Control de producción (jefaturas)', 'usuarios':'Administración de usuarios'
}

class Usuario(db.Model):
    id=db.Column(db.Integer, primary_key=True); username=db.Column(db.String(60), unique=True, nullable=False)
    nombre=db.Column(db.String(120), nullable=False); cargo=db.Column(db.String(100)); rut=db.Column(db.String(20)); email=db.Column(db.String(140)); sigla=db.Column(db.String(10), unique=True); color=db.Column(db.String(20), unique=True); password_hash=db.Column(db.String(255), nullable=False)
    disponible_planificacion=db.Column(db.Boolean, default=False); rol=db.Column(db.String(30), default='Usuario'); permisos_json=db.Column(db.Text, default='[]'); activo=db.Column(db.Boolean, default=True)
    creado=db.Column(db.DateTime, default=datetime.utcnow); ultimo_acceso=db.Column(db.DateTime)
    def set_password(self,p): self.password_hash=generate_password_hash(p)
    def check_password(self,p): return check_password_hash(self.password_hash,p)
    @property
    def es_admin(self):
        return (self.rol or '').strip().lower() in ('administrador','admin','administrator')
    @property
    def permisos(self):
        if self.es_admin: return list(PERMISOS.keys())
        if (self.rol or '').strip().casefold()=='supervisor de producción': return ['ingreso','diagnostico','planificacion','produccion']
        try: return json.loads(self.permisos_json or '[]')
        except: return []
    def puede(self,p):
        # Compatibilidad con permisos de versiones anteriores.
        if self.es_admin: return True
        if (self.rol or '').strip().casefold()=='supervisor de producción': return p in ('ingreso','diagnostico','planificacion','produccion')
        permisos=set(self.permisos)
        aliases={'usuarios':{'administracion'}, 'clientes':{'maestros'}}
        return p in permisos or bool(aliases.get(p,set()) & permisos)

class ReparacionMaestra(db.Model):
    id=db.Column(db.Integer, primary_key=True)
    componente_modelo_id=db.Column(db.Integer, db.ForeignKey('componente_modelo.id'), nullable=False, index=True)
    nombre=db.Column(db.String(180), nullable=False)
    descripcion=db.Column(db.Text)
    proceso=db.Column(db.String(100))
    hh_estimadas=db.Column(db.Float, default=0)
    recurso=db.Column(db.String(100))
    maquina=db.Column(db.String(120))
    secuencia=db.Column(db.Integer, default=0)
    activo=db.Column(db.Boolean, default=True)
    componente_modelo=db.relationship('ComponenteModelo', backref=db.backref('reparaciones', cascade='all, delete-orphan', lazy=True))

class OTReparacionSeleccionada(db.Model):
    id=db.Column(db.Integer,primary_key=True)
    ot_id=db.Column(db.Integer,db.ForeignKey('ot.id'),nullable=False,index=True)
    ot_componente_id=db.Column(db.Integer,db.ForeignKey('ot_componente.id'),nullable=False,index=True)
    reparacion_maestra_id=db.Column(db.Integer,db.ForeignKey('reparacion_maestra.id'),nullable=False)
    __table_args__=(db.UniqueConstraint('ot_componente_id','reparacion_maestra_id',name='uq_ot_reparacion'),)

class PlanNoAplica(db.Model):
    id=db.Column(db.Integer, primary_key=True)
    ot_id=db.Column(db.Integer, db.ForeignKey('ot.id'), nullable=False, index=True)
    ot_componente_id=db.Column(db.Integer, db.ForeignKey('ot_componente.id'), nullable=False, index=True)
    reparacion_maestra_id=db.Column(db.Integer, db.ForeignKey('reparacion_maestra.id'), nullable=False)
    fecha=db.Column(db.DateTime, default=datetime.utcnow)
    usuario=db.Column(db.String(120))
    __table_args__=(db.UniqueConstraint('ot_componente_id','reparacion_maestra_id',name='uq_plan_no_aplica'),)

class PlanActividad(db.Model):
    id=db.Column(db.Integer, primary_key=True)
    ot_id=db.Column(db.Integer, db.ForeignKey('ot.id'), nullable=False, index=True)
    ot_componente_id=db.Column(db.Integer, db.ForeignKey('ot_componente.id'))
    reparacion_maestra_id=db.Column(db.Integer, db.ForeignKey('reparacion_maestra.id'))
    usuario_id=db.Column(db.Integer, db.ForeignKey('usuario.id'), nullable=False)
    tarea=db.Column(db.String(200), nullable=False)
    fecha_inicio=db.Column(db.Date, nullable=False)
    fecha_fin=db.Column(db.Date, nullable=False)
    horas_dia=db.Column(db.Float, default=8)
    estado=db.Column(db.String(30), default='Planificada')
    observacion=db.Column(db.Text)
    ot=db.relationship('OT')
    componente=db.relationship('OTComponente')
    reparacion=db.relationship('ReparacionMaestra')
    responsable=db.relationship('Usuario')


class PlanLiberacion(db.Model):
    """Autorización individual, de un solo uso, para editar una asignación."""
    id=db.Column(db.Integer,primary_key=True)
    asignacion_id=db.Column(db.Integer,db.ForeignKey('plan_actividad.id'),nullable=False,unique=True,index=True)
    liberado_por_id=db.Column(db.Integer,db.ForeignKey('usuario.id'),nullable=False)
    liberado_fecha=db.Column(db.DateTime,nullable=False,default=datetime.utcnow)
    consumido_por_id=db.Column(db.Integer,db.ForeignKey('usuario.id'))
    consumido_fecha=db.Column(db.DateTime)


class ProduccionEtapa(db.Model):
    id=db.Column(db.Integer,primary_key=True)
    ot_id=db.Column(db.Integer,db.ForeignKey('ot.id'),nullable=False,index=True)
    ot_componente_id=db.Column(db.Integer,db.ForeignKey('ot_componente.id'),nullable=False,index=True)
    reparacion_maestra_id=db.Column(db.Integer,db.ForeignKey('reparacion_maestra.id'),nullable=False,index=True)
    estado=db.Column(db.String(25),nullable=False,default='Pendiente')
    observacion=db.Column(db.Text)
    terminado_por_id=db.Column(db.Integer,db.ForeignKey('usuario.id'))
    terminado_fecha=db.Column(db.DateTime)
    __table_args__=(db.UniqueConstraint('ot_componente_id','reparacion_maestra_id',name='uq_produccion_etapa'),)

class Auditoria(db.Model):
    id=db.Column(db.Integer, primary_key=True); fecha=db.Column(db.DateTime, default=datetime.utcnow)
    usuario=db.Column(db.String(120)); accion=db.Column(db.String(200)); detalle=db.Column(db.Text)

class Cliente(db.Model):
    id=db.Column(db.Integer, primary_key=True)
    nombre=db.Column(db.String(120), unique=True, nullable=False) # razón social / nombre histórico
    nombre_comercial=db.Column(db.String(120))
    rut=db.Column(db.String(20), unique=True)
    direccion=db.Column(db.String(200)); comuna=db.Column(db.String(100)); ciudad=db.Column(db.String(100)); giro=db.Column(db.String(180))
    telefono=db.Column(db.String(60)); email=db.Column(db.String(140)); activo=db.Column(db.Boolean, default=True)
    contactos=db.relationship('ContactoCliente', backref='empresa', cascade='all, delete-orphan', lazy=True)

class ContactoCliente(db.Model):
    id=db.Column(db.Integer, primary_key=True); cliente_id=db.Column(db.Integer, db.ForeignKey('cliente.id'), nullable=False)
    nombre=db.Column(db.String(140), nullable=False); cargo=db.Column(db.String(120)); email=db.Column(db.String(140)); telefono=db.Column(db.String(60))
    recibe_cotizacion=db.Column(db.Boolean, default=True); puede_aprobar=db.Column(db.Boolean, default=False)
    recibe_diagnostico=db.Column(db.Boolean, default=True); recibe_informe_final=db.Column(db.Boolean, default=True); activo=db.Column(db.Boolean, default=True)
class ModeloEquipo(db.Model):
    id=db.Column(db.Integer, primary_key=True); tipo=db.Column(db.String(80), nullable=False); marca=db.Column(db.String(80), nullable=False); modelo=db.Column(db.String(80), nullable=False); activo=db.Column(db.Boolean, default=True)
    componentes=db.relationship('ComponenteModelo', backref='modelo_equipo', cascade='all, delete-orphan', lazy=True)
    __table_args__=(db.UniqueConstraint('tipo','marca','modelo',name='uq_equipo_modelo'),)
class ComponenteModelo(db.Model):
    id=db.Column(db.Integer, primary_key=True); modelo_id=db.Column(db.Integer, db.ForeignKey('modelo_equipo.id'), nullable=False); nombre=db.Column(db.String(100), nullable=False)
    codigo=db.Column(db.String(80)); categoria=db.Column(db.String(40), default='Componente reparable'); material_default=db.Column(db.String(80)); activo=db.Column(db.Boolean, default=True)

class OTComponente(db.Model):
    id=db.Column(db.Integer, primary_key=True); ot_id=db.Column(db.Integer, db.ForeignKey('ot.id'), nullable=False); componente_modelo_id=db.Column(db.Integer, db.ForeignKey('componente_modelo.id'))
    nombre=db.Column(db.String(120), nullable=False); codigo=db.Column(db.String(80)); categoria=db.Column(db.String(40), default='Componente reparable')
    condicion_ingreso=db.Column(db.String(30), default='No recibido'); resolucion=db.Column(db.String(40), default='Sin intervención')
    material=db.Column(db.String(80)); ubicacion=db.Column(db.String(120)); observacion=db.Column(db.Text)
    componente_modelo=db.relationship('ComponenteModelo')

class DiagnosticoHallazgo(db.Model):
    id=db.Column(db.Integer, primary_key=True)
    ot_id=db.Column(db.Integer, db.ForeignKey('ot.id'), nullable=False, index=True)
    ot_componente_id=db.Column(db.Integer, db.ForeignKey('ot_componente.id'))
    zona=db.Column(db.String(160))
    condicion=db.Column(db.Text, nullable=False)
    trabajo_recomendado=db.Column(db.Text)
    observacion=db.Column(db.Text)
    estado=db.Column(db.String(40), default='Pendiente')
    foto_nombre=db.Column(db.String(180))
    foto_mime=db.Column(db.String(60))
    foto_datos=db.Column(db.LargeBinary)
    creado=db.Column(db.DateTime, default=datetime.utcnow)
    creado_por=db.Column(db.String(120))
    componente=db.relationship('OTComponente')

class DiagnosticoMarca(db.Model):
    __tablename__='diagnostico_marca'
    id=db.Column(db.Integer,primary_key=True)
    hallazgo_id=db.Column(db.Integer,db.ForeignKey('diagnostico_hallazgo.id'),unique=True,nullable=False)
    marcas_json=db.Column(db.Text,default='[]')

class DiagnosticoFoto(db.Model):
    __tablename__='diagnostico_foto'
    id=db.Column(db.Integer,primary_key=True)
    hallazgo_id=db.Column(db.Integer,db.ForeignKey('diagnostico_hallazgo.id'),nullable=False,index=True)
    nombre=db.Column(db.String(180)); mime=db.Column(db.String(60)); datos=db.Column(db.LargeBinary,nullable=False)
    marcas_json=db.Column(db.Text,default='[]')
    creado=db.Column(db.DateTime,default=datetime.utcnow)
    hallazgo=db.relationship('DiagnosticoHallazgo',backref=db.backref('fotos_adicionales',cascade='all, delete-orphan',lazy=True,order_by='DiagnosticoFoto.id'))

class Cotizacion(db.Model):
    id=db.Column(db.Integer, primary_key=True)
    numero=db.Column(db.String(30), unique=True, nullable=False)
    fecha=db.Column(db.DateTime, default=datetime.utcnow)
    validez_hasta=db.Column(db.Date)
    ot_id=db.Column(db.Integer, db.ForeignKey('ot.id'), nullable=False)
    cliente=db.Column(db.String(120), nullable=False)
    contacto=db.Column(db.String(120))
    moneda=db.Column(db.String(10), default='CLP')
    plazo_pago=db.Column(db.String(80), default='30 días')
    tiempo_entrega=db.Column(db.String(120))
    plazo_ejecucion_cantidad=db.Column(db.Integer)
    plazo_ejecucion_unidad=db.Column(db.String(30), default='Días hábiles')
    inicio_plazo=db.Column(db.String(120), default='Recepción de OC / aprobación')
    fecha_aprobacion=db.Column(db.DateTime)
    fecha_comprometida=db.Column(db.Date)
    referencia=db.Column(db.String(120))
    notas=db.Column(db.Text)
    estado=db.Column(db.String(30), default='Presupuesto')
    descuento_pct=db.Column(db.Float, default=0)
    anulada_motivo=db.Column(db.Text)
    anulada_fecha=db.Column(db.DateTime)
    anulada_por=db.Column(db.String(120))
    creada_por=db.Column(db.String(120))
    lineas=db.relationship('LineaCotizacion', backref='cotizacion', cascade='all, delete-orphan', lazy=True)
    @property
    def subtotal(self): return sum((l.cantidad or 0)*(l.precio_unitario or 0) for l in self.lineas)
    @property
    def descuento(self): return round(self.subtotal * max(0, min(self.descuento_pct or 0, 100)) / 100)
    @property
    def neto(self): return self.subtotal-self.descuento
    @property
    def iva(self): return round(self.neto*0.19)
    @property
    def total(self): return self.neto+self.iva

class LineaCotizacion(db.Model):
    id=db.Column(db.Integer, primary_key=True)
    cotizacion_id=db.Column(db.Integer, db.ForeignKey('cotizacion.id'), nullable=False)
    descripcion=db.Column(db.String(250), nullable=False)
    cantidad=db.Column(db.Float, default=1)
    precio_unitario=db.Column(db.Integer, default=0)
    orden=db.Column(db.Integer, default=0)

class OT(db.Model):
    id=db.Column(db.Integer, primary_key=True); modelo_id=db.Column(db.Integer, db.ForeignKey('modelo_equipo.id')); numero=db.Column(db.String(30), unique=True, nullable=False); fecha=db.Column(db.DateTime, default=datetime.utcnow)
    cliente=db.Column(db.String(120), nullable=False); equipo=db.Column(db.String(120), nullable=False); componente=db.Column(db.String(500), nullable=False)
    serie=db.Column(db.String(80)); guia=db.Column(db.String(80)); estado=db.Column(db.String(40), default='Diagnóstico'); diagnostico=db.Column(db.Text)
    cotizacion=db.Column(db.String(80)); aprobacion=db.Column(db.String(30), default='Pendiente'); reparacion=db.Column(db.Text); control_calidad=db.Column(db.Text)
    checklist=db.Column(db.String(120)); embalaje=db.Column(db.Text); despacho=db.Column(db.Text)
    fecha_comprometida=db.Column(db.Date)
    fecha_entrega_faena=db.Column(db.Date)
    recibido_por=db.Column(db.String(160))
    revisado_faena_por=db.Column(db.String(160))
    fecha_envio_cotizacion=db.Column(db.Date)
    enviado_por=db.Column(db.String(160))
    enviado_a=db.Column(db.String(160))
    recibido_bodega_por=db.Column(db.String(160))
    observaciones_admin=db.Column(db.Text)
    modelo_equipo=db.relationship('ModeloEquipo')
    componentes_detalle=db.relationship('OTComponente', backref='ot', cascade='all, delete-orphan', lazy=True)
    hallazgos=db.relationship('DiagnosticoHallazgo', backref='ot', cascade='all, delete-orphan', lazy=True, order_by='DiagnosticoHallazgo.id')
    @property
    def avance(self):
        estados=['Diagnóstico','Cotización enviada','Aprobada / Reparación','Control de calidad','Checklist / Embalaje','Despachada','Cerrada']
        try:return int((estados.index(self.estado)+1)/len(estados)*100)
        except:return 10

def usuario_actual(): return db.session.get(Usuario, session.get('user_id')) if session.get('user_id') else None

def es_supervisor():
    u=usuario_actual()
    return bool(u and not u.es_admin and (u.rol or '').strip().casefold()=='supervisor de producción')

# Seguridad del servidor: listado explícito de rutas permitidas para supervisión.
# Todo endpoint nuevo queda BLOQUEADO por defecto hasta revisión.
SUPERVISOR_GET={'index','control_general_ot','nueva_ot','ver_ot','diagnostico_ot',
    'diagnostico_foto','diagnostico_foto_adicional','diagnostico_pdf',
    'diagnostico_editar','produccion','planificacion_gantt','planificacion_gantt_pdf',
    'componentes_modelo','oti_lista','oti_seleccionar','oti_editar','oti_pdf','logout','login','static'}
SUPERVISOR_POST={'nueva_ot','diagnostico_ot','diagnostico_editar',
    'diagnostico_marcas','diagnostico_foto_marcas','diagnostico_foto_eliminar',
    'produccion_terminar','oti_nueva','oti_editar','api_gantt_rapido'}

@app.before_request
def supervisores_rutas_seguras():
    if not es_supervisor(): return
    endpoint=request.endpoint or ''
    permitidos=SUPERVISOR_GET if request.method in ('GET','HEAD') else SUPERVISOR_POST
    if endpoint not in permitidos: abort(403)
    # La OT es de solo lectura; la recepción se realiza mediante Nueva OT.
    if endpoint=='ver_ot' and request.method!='GET': abort(403)

@app.context_processor
def inject_user(): return dict(usuario_actual=usuario_actual(), es_supervisor=es_supervisor(), PERMISOS=PERMISOS, cotizacion_activa_ot=lambda ot_id: Cotizacion.query.filter(Cotizacion.ot_id==ot_id, Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).first())

def login_required(fn):
    @wraps(fn)
    def w(*a,**k):
        if Usuario.query.count()==0: return redirect(url_for('setup'))
        u=usuario_actual()
        if not u or not u.activo: session.clear(); return redirect(url_for('login'))
        return fn(*a,**k)
    return w

def permiso_required(p):
    def deco(fn):
        @wraps(fn)
        @login_required
        def w(*a,**k):
            if not usuario_actual().puede(p): abort(403)
            return fn(*a,**k)
        return w
    return deco

def audit(accion, detalle=''):
    u=usuario_actual(); db.session.add(Auditoria(usuario=u.nombre if u else 'Sistema', accion=accion, detalle=detalle)); db.session.commit()

def normalizar_rut(rut):
    raw=''.join(ch for ch in (rut or '').upper() if ch.isdigit() or ch=='K')
    if len(raw)<2: return ''
    return raw[:-1]+'-'+raw[-1]

def rut_valido(rut):
    r=normalizar_rut(rut)
    if '-' not in r: return False
    cuerpo,dv=r.split('-')
    if not cuerpo.isdigit() or dv not in '0123456789K': return False
    suma=0; factor=2
    for d in reversed(cuerpo):
        suma += int(d)*factor; factor = 2 if factor==7 else factor+1
    x=11-(suma%11); esperado='0' if x==11 else 'K' if x==10 else str(x)
    return dv==esperado

def formatear_rut(rut):
    r=normalizar_rut(rut)
    if '-' not in r: return rut or ''
    cuerpo,dv=r.split('-'); partes=[]
    while cuerpo: partes.insert(0,cuerpo[-3:]); cuerpo=cuerpo[:-3]
    return '.'.join(partes)+'-'+dv

def seed_maestros():
    if Cliente.query.count()==0:
        for n in ['SQM Nueva Victoria','SQM Salar','Cliente de prueba']: db.session.add(Cliente(nombre=n))
    if ModeloEquipo.query.count()==0:
        m=ModeloEquipo(tipo='Bomba centrífuga',marca='Vogel',modelo='P204/5')
        m.componentes=[ComponenteModelo(nombre=n) for n in ['Motor','Base','Acoplamiento','Tapa rodamiento lado transmisión','Tapa rodamiento lado impulsión','Cuerpo de rodamientos','Eje','Impulsor','Tuerca impulsor','Voluta','Descarga','Válvula check','Rodamiento lado transmisión','Rodamiento lado impulsión','Sello laberinto / Retén','Empaquetadura','O-rings']]; db.session.add(m)
    db.session.commit()

def sumar_plazo(fecha_base, cantidad, unidad):
    if not fecha_base or not cantidad: return None
    cantidad=int(cantidad)
    if unidad=='Semanas': return fecha_base + timedelta(weeks=cantidad)
    if unidad=='Días corridos': return fecha_base + timedelta(days=cantidad)
    # Días hábiles: lunes a viernes; feriados se podrán incorporar en una versión posterior.
    d=fecha_base; agregados=0
    while agregados<cantidad:
        d += timedelta(days=1)
        if d.weekday()<5: agregados += 1
    return d

def asegurar_columnas():
    # Migración ligera para instalaciones V4 existentes: create_all no agrega columnas nuevas.
    from sqlalchemy import inspect, text
    insp=inspect(db.engine)
    if 'usuario' in insp.get_table_names():
        cols={c['name'] for c in insp.get_columns('usuario')}
        defs={'rut':'VARCHAR(20)','email':'VARCHAR(140)','sigla':'VARCHAR(10)','color':'VARCHAR(20)','disponible_planificacion':'BOOLEAN DEFAULT FALSE'}
        for n,t in defs.items():
            if n not in cols: db.session.execute(text(f'ALTER TABLE usuario ADD COLUMN {n} {t}'))
        try: db.session.execute(text('CREATE UNIQUE INDEX IF NOT EXISTS uq_usuario_sigla ON usuario (sigla) WHERE sigla IS NOT NULL'))
        except: pass
        try: db.session.execute(text('CREATE UNIQUE INDEX IF NOT EXISTS uq_usuario_color ON usuario (color) WHERE color IS NOT NULL'))
        except: pass
    if 'cotizacion' in insp.get_table_names():
        cols={c['name'] for c in insp.get_columns('cotizacion')}
        defs={
          'plazo_ejecucion_cantidad':'INTEGER', 'plazo_ejecucion_unidad':'VARCHAR(30)',
          'inicio_plazo':'VARCHAR(120)', 'fecha_aprobacion':'TIMESTAMP', 'fecha_comprometida':'DATE',
          'descuento_pct':'FLOAT', 'anulada_motivo':'TEXT', 'anulada_fecha':'TIMESTAMP', 'anulada_por':'VARCHAR(120)'}
        for n,t in defs.items():
            if n not in cols: db.session.execute(text(f'ALTER TABLE cotizacion ADD COLUMN {n} {t}'))
    if 'ot' in insp.get_table_names():
        cols={c['name'] for c in insp.get_columns('ot')}
        if 'fecha_comprometida' not in cols: db.session.execute(text('ALTER TABLE ot ADD COLUMN fecha_comprometida DATE'))
        if 'modelo_id' not in cols: db.session.execute(text('ALTER TABLE ot ADD COLUMN modelo_id INTEGER'))
        for n,t in {'fecha_entrega_faena':'DATE','recibido_por':'VARCHAR(160)','revisado_faena_por':'VARCHAR(160)','fecha_envio_cotizacion':'DATE','enviado_por':'VARCHAR(160)','enviado_a':'VARCHAR(160)','recibido_bodega_por':'VARCHAR(160)','observaciones_admin':'TEXT'}.items():
            if n not in cols: db.session.execute(text(f'ALTER TABLE ot ADD COLUMN {n} {t}'))
    if 'cliente' in insp.get_table_names():
        cols={c['name'] for c in insp.get_columns('cliente')}
        defs={'nombre_comercial':'VARCHAR(120)','rut':'VARCHAR(20)','direccion':'VARCHAR(200)','comuna':'VARCHAR(100)','ciudad':'VARCHAR(100)','giro':'VARCHAR(180)','telefono':'VARCHAR(60)','email':'VARCHAR(140)'}
        for n,t in defs.items():
            if n not in cols: db.session.execute(text(f'ALTER TABLE cliente ADD COLUMN {n} {t}'))
        try: db.session.execute(text('CREATE UNIQUE INDEX IF NOT EXISTS uq_cliente_rut ON cliente (rut) WHERE rut IS NOT NULL'))
        except: pass
    if 'modelo_equipo' in insp.get_table_names():
        cols={c['name'] for c in insp.get_columns('modelo_equipo')}
        if 'materialidad' not in cols: db.session.execute(text('ALTER TABLE modelo_equipo ADD COLUMN materialidad VARCHAR(120)'))
    if 'componente_modelo' in insp.get_table_names():
        cols={c['name'] for c in insp.get_columns('componente_modelo')}
        defs={'codigo':'VARCHAR(80)','categoria':'VARCHAR(40)','material_default':'VARCHAR(80)','activo':'BOOLEAN DEFAULT TRUE'}
        for n,t in defs.items():
            if n not in cols: db.session.execute(text(f'ALTER TABLE componente_modelo ADD COLUMN {n} {t}'))
    db.session.commit()

def next_quote_number():
    y=datetime.now().year; last=Cotizacion.query.filter(Cotizacion.numero.like(f'COT-{y}-%')).order_by(Cotizacion.id.desc()).first(); seq=int(last.numero.split('-')[-1])+1 if last else 1
    return f'COT-{y}-{seq:04d}'

def next_number():
    y=datetime.now().year; last=OT.query.filter(OT.numero.like(f'OT-{y}-%')).order_by(OT.id.desc()).first(); seq=int(last.numero.split('-')[-1])+1 if last else 1
    return f'OT-{y}-{seq:04d}'

def preparar_foto(upload):
    if not upload or not upload.filename: return None, None, None
    try:
        img=PILImage.open(upload.stream)
        img=ImageOps.exif_transpose(img).convert('RGB')
        img.thumbnail((1600,1600))
        out=BytesIO(); img.save(out,format='JPEG',quality=82,optimize=True)
        return upload.filename[:180], 'image/jpeg', out.getvalue()
    except Exception:
        return None, None, None

@app.route('/setup',methods=['GET','POST'])
def setup():
    if Usuario.query.count()>0: return redirect(url_for('login'))
    if request.method=='POST':
        if len(request.form['password'])<8: flash('La contraseña debe tener al menos 8 caracteres.'); return render_template('setup.html')
        u=Usuario(username=request.form['username'].strip().lower(), nombre=request.form['nombre'].strip(), cargo='Administrador', rol='Administrador'); u.set_password(request.form['password'])
        db.session.add(u); db.session.commit(); session['user_id']=u.id; audit('Creación administrador inicial'); return redirect(url_for('index'))
    return render_template('setup.html')

@app.route('/login',methods=['GET','POST'])
def login():
    if Usuario.query.count()==0:return redirect(url_for('setup'))
    if request.method=='POST':
        u=Usuario.query.filter_by(username=request.form['username'].strip().lower()).first()
        if u and u.activo and u.check_password(request.form['password']):
            session.clear(); session['user_id']=u.id; u.ultimo_acceso=datetime.utcnow(); db.session.commit(); audit('Inicio de sesión'); return redirect(url_for('index'))
        flash('Usuario o contraseña incorrectos.')
    return render_template('login.html')
@app.route('/logout')
def logout(): session.clear(); return redirect(url_for('login'))

def listado_control_ot():
    """Listado administrativo de OT con los mismos filtros para pantalla y Excel."""
    ots=OT.query.order_by(OT.id.desc()).all()
    ids=[o.id for o in ots]
    cotizaciones=Cotizacion.query.filter(Cotizacion.ot_id.in_(ids),Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).all() if ids else []
    por_ot={}
    for q in cotizaciones:
        por_ot.setdefault(q.ot_id,q)
    planes=set(x[0] for x in db.session.query(PlanActividad.ot_id).filter(PlanActividad.ot_id.in_(ids)).distinct().all()) if ids else set()
    filas=[]
    for o in ots:
        q=por_ot.get(o.id)
        diag=bool((o.diagnostico or '').strip() or o.hallazgos)
        qe=(q.estado or '').casefold() if q else ''
        aprobado=qe=='aprobada'
        if o.fecha_entrega_faena and (o.recibido_por or '').strip(): estado='Entregada'
        elif (o.despacho or '').strip() and (o.guia or '').strip(): estado='Despachada'
        elif (o.control_calidad or '').strip() and (o.checklist or '').strip(): estado='Lista para despacho'
        elif (o.reparacion or '').strip(): estado='En reparación'
        elif o.id in planes and aprobado: estado='Planificada'
        elif qe=='rechazada': estado='Cotización rechazada'
        elif aprobado: estado='Aprobada'
        elif qe=='presupuesto enviado': estado='Cotización enviada'
        elif q: estado='Cotizada'
        elif diag: estado='Diagnosticada'
        else: estado='En diagnóstico'
        filas.append(dict(ot=o,cot=q,estado=estado,precio=q.neto if q else None))
    buscar=(request.args.get('buscar') or '').strip().casefold()
    filtro=(request.args.get('estado') or '').strip()
    if buscar:
        filas=[r for r in filas if buscar in ' '.join(str(x or '') for x in (r['ot'].numero,r['ot'].cliente,r['ot'].equipo,r['ot'].componente,r['ot'].serie,r['ot'].ot_cliente if hasattr(r['ot'],'ot_cliente') else '',r['cot'].numero if r['cot'] else '')).casefold()]
    if filtro:
        filas=[r for r in filas if r['estado']==filtro]
    return filas

@app.route('/')
@login_required
def index():
    filas=listado_control_ot()
    if es_supervisor():
        for f in filas: f['precio']=None; f['cot']=None
    return render_template('index.html',filas=filas,buscar=request.args.get('buscar',''),filtro=request.args.get('estado',''))

@app.route('/ot')
@login_required
def control_general_ot():
    filas=listado_control_ot()
    if es_supervisor():
        for f in filas: f['precio']=None; f['cot']=None
    return render_template('index.html',filas=filas,buscar=request.args.get('buscar',''),filtro=request.args.get('estado',''),solo_ot=True)

@app.route('/ot/exportar-excel')
@permiso_required('ingreso')
def exportar_ot_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook(); ws=wb.active; ws.title='Control OT'
    columnas=['OT INGEPRO','Fecha ingreso','Cliente','Equipo','Componente','Serie','OT cliente','Guía despacho','Cotización','Fecha cotización','Precio neto CLP','Estado','Fecha envío cotización','Enviado por','Enviado a','Fecha entrega faena','Recibido por','Revisado faena INGEPRO']
    ws.append(columnas)
    for r in listado_control_ot():
        o=r['ot']; q=r['cot']
        ws.append([o.numero,o.fecha.date() if o.fecha else None,o.cliente,o.equipo,o.componente,o.serie,getattr(o,'ot_cliente',None),o.guia,q.numero if q else None,q.fecha.date() if q and q.fecha else None,r['precio'],r['estado'],o.fecha_envio_cotizacion,o.enviado_por,o.enviado_a,o.fecha_entrega_faena,o.recibido_por,o.revisado_faena_por])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True)
    for col in ws.columns:
        letter=col[0].column_letter
        ws.column_dimensions[letter].width=min(38,max(15,max(len(str(c.value or '')) for c in col[:100])+2))
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    buf=BytesIO();wb.save(buf);buf.seek(0)
    return send_file(buf,as_attachment=True,download_name='INGEPRO_Control_OT.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/usuarios')
@permiso_required('usuarios')
def usuarios(): return render_template('usuarios.html',usuarios=Usuario.query.order_by(Usuario.nombre).all())
@app.route('/usuarios/nuevo',methods=['GET','POST'])
@permiso_required('usuarios')
def usuario_nuevo():
    if request.method=='POST':
        username=request.form['username'].strip().lower()
        if Usuario.query.filter_by(username=username).first(): flash('Ese usuario ya existe.'); return render_template('usuario_form.html',permisos=PERMISOS)
        if len(request.form['password'])<8: flash('La contraseña debe tener al menos 8 caracteres.'); return render_template('usuario_form.html',permisos=PERMISOS)
        sigla=request.form.get('sigla','').strip().upper() or None; color=request.form.get('color','').strip() or None
        if sigla and Usuario.query.filter_by(sigla=sigla).first(): flash('La sigla ya está asignada a otro usuario.'); return render_template('usuario_form.html',permisos=PERMISOS,usuario=None)
        if color and Usuario.query.filter_by(color=color).first(): flash('El color ya está asignado a otro usuario.'); return render_template('usuario_form.html',permisos=PERMISOS,usuario=None)
        u=Usuario(username=username,nombre=request.form['nombre'].strip(),cargo=request.form.get('cargo','').strip(),rut=request.form.get('rut','').strip(),email=request.form.get('email','').strip(),sigla=sigla,color=color,rol=request.form.get('rol','Usuario'),permisos_json=json.dumps(['ingreso','diagnostico','planificacion','produccion'] if request.form.get('rol')=='Supervisor de Producción' else request.form.getlist('permisos')),disponible_planificacion=('disponible_planificacion' in request.form)); u.set_password(request.form['password'])
        db.session.add(u); db.session.commit(); audit('Usuario creado',u.username); flash('Usuario creado correctamente.'); return redirect(url_for('usuarios'))
    return render_template('usuario_form.html',permisos=PERMISOS,usuario=None)

@app.route('/usuarios/<int:uid>/editar',methods=['GET','POST'])
@permiso_required('usuarios')
def usuario_editar(uid):
    u=db.session.get(Usuario,uid) or abort(404)
    if request.method=='POST':
        sigla=request.form.get('sigla','').strip().upper() or None; color=request.form.get('color','').strip() or None
        if sigla and Usuario.query.filter(Usuario.sigla==sigla,Usuario.id!=u.id).first(): flash('La sigla ya está asignada a otro usuario.'); return render_template('usuario_form.html',permisos=PERMISOS,usuario=u)
        if color and Usuario.query.filter(Usuario.color==color,Usuario.id!=u.id).first(): flash('El color ya está asignado a otro usuario.'); return render_template('usuario_form.html',permisos=PERMISOS,usuario=u)
        u.nombre=request.form['nombre'].strip(); u.cargo=request.form.get('cargo','').strip(); u.rut=request.form.get('rut','').strip(); u.email=request.form.get('email','').strip(); u.sigla=sigla; u.color=color; u.rol=request.form.get('rol','Usuario'); u.permisos_json=json.dumps(['ingreso','diagnostico','planificacion','produccion'] if request.form.get('rol')=='Supervisor de Producción' else request.form.getlist('permisos')); u.disponible_planificacion=('disponible_planificacion' in request.form)
        if request.form.get('password'): u.set_password(request.form['password'])
        db.session.commit(); audit('Usuario modificado',u.username); flash('Usuario actualizado.'); return redirect(url_for('usuarios'))
    return render_template('usuario_form.html',permisos=PERMISOS,usuario=u)

@app.route('/usuarios/<int:uid>/estado',methods=['POST'])
@permiso_required('usuarios')
def usuario_estado(uid):
    u=db.session.get(Usuario,uid) or abort(404)
    if u.id==usuario_actual().id: flash('No puedes desactivar tu propia cuenta.'); return redirect(url_for('usuarios'))
    u.activo=not u.activo; db.session.commit(); audit('Estado de usuario modificado',f'{u.username}: {u.activo}'); return redirect(url_for('usuarios'))

@app.route('/api/modelo/<int:modelo_id>/componentes')
@login_required
def componentes_modelo(modelo_id):
    m=db.session.get(ModeloEquipo,modelo_id) or abort(404); return jsonify([{'id':c.id,'nombre':c.nombre,'codigo':c.codigo or '','categoria':c.categoria or 'Componente reparable','material':c.material_default or ''} for c in m.componentes if c.activo is not False])
@app.route('/ot/nueva',methods=['GET','POST'])
@permiso_required('ingreso')
def nueva_ot():
    clientes=Cliente.query.filter_by(activo=True).order_by(Cliente.nombre).all(); modelos=ModeloEquipo.query.filter_by(activo=True).order_by(ModeloEquipo.tipo,ModeloEquipo.marca,ModeloEquipo.modelo).all()
    if request.method=='POST':
        cliente=db.session.get(Cliente,int(request.form['cliente_id'])) or abort(404); modelo=db.session.get(ModeloEquipo,int(request.form['modelo_id'])) or abort(404)
        recibidos=set(request.form.getlist('componentes_recibidos'))
        ot=OT(numero=next_number(),modelo_id=modelo.id,cliente=cliente.nombre,equipo=f'{modelo.tipo} {modelo.marca} {modelo.modelo}',componente='Pendiente detalle',serie=request.form.get('serie'),guia=request.form.get('guia'),recibido_bodega_por=(usuario_actual().nombre if usuario_actual() else None))
        db.session.add(ot); db.session.flush()
        nombres=[]
        for c in modelo.componentes:
            if c.activo is False: continue
            cond='Recibido' if str(c.id) in recibidos else 'No recibido'
            if cond=='Recibido': nombres.append(c.nombre)
            db.session.add(OTComponente(ot_id=ot.id,componente_modelo_id=c.id,nombre=c.nombre,codigo=c.codigo,categoria=c.categoria or 'Componente reparable',condicion_ingreso=cond,material=c.material_default))
        ot.componente=', '.join(nombres) if nombres else 'Sin componentes marcados como recibidos'
        db.session.commit(); audit('OT creada',ot.numero); flash(f'{ot.numero} creada correctamente'); return redirect(url_for('ver_ot',ot_id=ot.id))
    # El formulario utiliza `catalogo|tojson`: pasar solo valores JSON serializables.
    catalogo=[]
    for m in modelos:
        materiales=sorted({(c.material_default or '').strip() for c in m.componentes if c.activo is not False and (c.material_default or '').strip()})
        materialidad=' / '.join(materiales)
        catalogo.append({'id':m.id,'tipo':m.tipo or '', 'marca':m.marca or '', 'modelo':m.modelo or '', 'materialidad':materialidad})
    return render_template('nueva.html',clientes=clientes,modelos=modelos,catalogo=catalogo)
def cargar_despiece_ot(ot, modelo):
    if ot.componentes_detalle:
        return 0
    ot.modelo_id=modelo.id
    creados=0
    for c in modelo.componentes:
        if c.activo is False: continue
        db.session.add(OTComponente(ot_id=ot.id,componente_modelo_id=c.id,nombre=c.nombre,codigo=c.codigo,categoria=c.categoria or 'Componente reparable',condicion_ingreso='No recibido',material=c.material_default))
        creados += 1
    return creados

def inferir_modelo_ot(ot):
    if ot.modelo_id:
        return db.session.get(ModeloEquipo,ot.modelo_id)
    texto=(ot.equipo or '').lower()
    candidatos=ModeloEquipo.query.filter_by(activo=True).all()
    for m in candidatos:
        if m.marca.lower() in texto and m.modelo.lower() in texto:
            return m
    return None

@app.route('/ot/<int:ot_id>/cargar-despiece',methods=['POST'])
@permiso_required('ingreso')
def ot_cargar_despiece(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    if ot.componentes_detalle:
        flash('Esta OT ya tiene un despiece cargado.'); return redirect(url_for('ver_ot',ot_id=ot.id))
    mid=request.form.get('modelo_id')
    modelo=db.session.get(ModeloEquipo,int(mid)) if mid else inferir_modelo_ot(ot)
    if not modelo:
        flash('Selecciona el modelo de bomba para cargar su despiece.'); return redirect(url_for('ver_ot',ot_id=ot.id))
    n=cargar_despiece_ot(ot,modelo)
    db.session.commit(); audit('Despiece cargado en OT',f'{ot.numero} / {modelo.marca} {modelo.modelo} / {n} componentes')
    flash(f'Despiece {modelo.marca} {modelo.modelo} cargado: {n} componentes. Marca ahora qué elementos fueron recibidos.')
    return redirect(url_for('ver_ot',ot_id=ot.id))

def seguimiento_ot(ot):
    """Estados derivados de registros reales; no se pueden marcar manualmente."""
    q=Cotizacion.query.filter(Cotizacion.ot_id==ot.id,Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).first()
    diag=bool((ot.diagnostico or '').strip() or ot.hallazgos)
    cot=bool(q)
    enviada=bool(q and (q.estado or '').strip().casefold() in ('presupuesto enviado','aprobada','rechazada'))
    aprobada=bool(q and (q.estado or '').strip().casefold()=='aprobada')
    rechazada=bool(q and (q.estado or '').strip().casefold()=='rechazada')
    actividades=PlanActividad.query.filter_by(ot_id=ot.id).first() if aprobada else None
    plan=bool(actividades)
    reparacion=bool((ot.reparacion or '').strip())
    calidad=bool((ot.control_calidad or '').strip() and (ot.checklist or '').strip())
    despacho=bool((ot.despacho or '').strip() and (ot.guia or '').strip())
    entrega=bool(ot.fecha_entrega_faena and (ot.recibido_por or '').strip())
    definiciones=[
      ('Ingreso','Equipo y OT cliente',True,'ver_ot','ingreso'),
      ('Diagnóstico','Informe técnico',diag,'diagnostico_ot','diagnostico'),
      ('Cotización','Precio y documento',cot,'ver_cotizacion' if q else 'cotizacion_nueva','cotizaciones'),
      ('Envío cotización','Fecha y destinatario',enviada,'ver_cotizacion' if q else 'cotizacion_nueva','cotizaciones'),
      ('Aprobación','Autorización cliente',aprobada,'ver_cotizacion' if q else 'cotizacion_nueva','cotizaciones'),
      ('Planificación','Gantt compacto',plan,'planificacion_gantt','planificacion'),
      ('Reparación','Trabajos por componente',reparacion,'ot_reparaciones','reparacion'),
      ('Control calidad','Inspección y pruebas',calidad,'ver_ot','calidad'),
      ('Despacho','Guía y salida',despacho,'ver_ot','despacho'),
      ('Entregada','Fecha y recepción',entrega,'ver_ot','despacho')]
    pasos=[]
    for i,(nombre,descripcion,hecho,endpoint,permiso) in enumerate(definiciones):
        previo=all(x['completo'] for x in pasos)
        estado='completa' if hecho and previo else ('pendiente' if previo and not rechazada else 'bloqueada')
        if i==0: estado='completa'
        if endpoint=='ver_cotizacion' and q: href=url_for(endpoint,qid=q.id)
        elif endpoint=='cotizacion_nueva': href=url_for(endpoint,ot_id=ot.id)
        elif endpoint in ('diagnostico_ot','ot_reparaciones'): href=url_for(endpoint,ot_id=ot.id)
        elif endpoint=='planificacion_gantt': href=url_for(endpoint,ot_id=ot.id)
        else: href=url_for('ver_ot',ot_id=ot.id, _anchor='ot-'+('entrega' if i==9 else 'control' if i==7 else 'despacho' if i==8 else 'datos'))
        pasos.append(dict(nombre=nombre,descripcion=descripcion,completo=bool(hecho and previo),estado=estado,href=href,permitido=usuario_actual().puede(permiso)))
    if rechazada: condicion='Cotización rechazada'
    elif entrega: condicion='Entregada'
    elif despacho: condicion='Despachada'
    elif calidad: condicion='Lista para despacho'
    elif reparacion: condicion='En reparación'
    elif plan: condicion='Planificada'
    elif aprobada: condicion='Aprobada'
    elif enviada: condicion='Cotización enviada'
    elif cot: condicion='Cotizada'
    elif diag: condicion='Diagnosticada'
    else: condicion='En diagnóstico'
    return pasos,condicion,q

@app.route('/ot/<int:ot_id>',methods=['GET','POST'])
@login_required
def ver_ot(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    if es_supervisor(): return render_template('ot_supervisor.html',ot=ot)
    if request.method=='POST':
        for f in ['diagnostico','reparacion','control_calidad','checklist','embalaje','despacho']:
            if f in request.form:setattr(ot,f,request.form.get(f))
        for c in ot.componentes_detalle:
            c.condicion_ingreso=request.form.get(f'condicion_{c.id}',c.condicion_ingreso)
            c.resolucion=request.form.get(f'resolucion_{c.id}',c.resolucion)
            c.material=request.form.get(f'material_{c.id}',c.material or '').strip()
            c.ubicacion=request.form.get(f'ubicacion_{c.id}',c.ubicacion or '').strip()
        # V7.37: campos editables de recepción y entrega; datos de creación inmutables.
        if 'fecha_entrega_faena' in request.form:
            raw=request.form.get('fecha_entrega_faena','').strip()
            try: ot.fecha_entrega_faena=date.fromisoformat(raw) if raw else None
            except ValueError:
                db.session.rollback(); flash('Fecha de entrega en faena inválida.'); return redirect(url_for('ver_ot',ot_id=ot.id))
        if 'guia' in request.form: ot.guia=request.form.get('guia','').strip()[:80]
        if 'enviado_por' in request.form: ot.enviado_por=request.form.get('enviado_por','').strip()[:160]
        if 'revisado_faena_por' in request.form: ot.revisado_faena_por=request.form.get('revisado_faena_por','').strip()[:160]
        if 'observaciones_admin' in request.form: ot.observaciones_admin=request.form.get('observaciones_admin','').strip()
        db.session.commit(); audit('OT actualizada',ot.numero); flash('OT actualizada'); return redirect(url_for('ver_ot',ot_id=ot.id))
    if not ot.componentes_detalle:
        modelo=inferir_modelo_ot(ot)
        if modelo:
            cargar_despiece_ot(ot,modelo); db.session.commit()
    modelos=ModeloEquipo.query.filter_by(activo=True).order_by(ModeloEquipo.tipo,ModeloEquipo.marca,ModeloEquipo.modelo).all()
    pasos,condicion,q=seguimiento_ot(ot)
    return render_template('ot.html',ot=ot,modelos=modelos,pasos=pasos,condicion=condicion,cot_activa=q)


@app.route('/ot/<int:ot_id>/agregar-componentes', methods=['POST'])
@permiso_required('ingreso')
def agregar_componentes_ot(ot_id):
    """Recibe componentes del catálogo en una OT existente sin alterar diagnósticos."""
    ot=db.session.get(OT, ot_id) or abort(404)
    modelo=ot.modelo_equipo or inferir_modelo_ot(ot)
    if not modelo:
        flash('Primero debes vincular un modelo a esta OT.')
        return redirect(url_for('ver_ot', ot_id=ot.id))
    valores=request.form.getlist('componentes_nuevos')
    ids=set()
    for valor in valores:
        try: ids.add(int(valor))
        except (TypeError, ValueError): abort(400)
    if not ids:
        flash('Selecciona al menos un componente.')
        return redirect(url_for('ver_ot', ot_id=ot.id))
    permitidos={c.id:c for c in modelo.componentes if c.activo is not False}
    if not ids.issubset(permitidos): abort(400)
    agregados=0
    reactivados=0
    for cid in sorted(ids):
        cm=permitidos[cid]
        existente=next((c for c in ot.componentes_detalle if c.componente_modelo_id==cid), None)
        if existente:
            if (existente.condicion_ingreso or '').strip().casefold()!='recibido':
                existente.condicion_ingreso='Recibido'
                reactivados+=1
        else:
            db.session.add(OTComponente(ot_id=ot.id, componente_modelo_id=cm.id,
                nombre=cm.nombre, codigo=cm.codigo,
                categoria=cm.categoria or 'Componente reparable',
                condicion_ingreso='Recibido', material=cm.material_default))
            agregados+=1
    db.session.flush()
    ot.componente=', '.join(c.nombre for c in ot.componentes_detalle
        if (c.condicion_ingreso or '').strip().casefold()=='recibido')
    db.session.commit()
    audit('Componentes recibidos agregados a OT', f'{ot.numero}: {agregados} nuevos, {reactivados} actualizados')
    flash(f'Componentes recibidos actualizados: {agregados} nuevos y {reactivados} existentes.')
    return redirect(url_for('ver_ot', ot_id=ot.id))

@app.route('/ot/<int:ot_id>/vincular-componente', methods=['POST'])
@permiso_required('ingreso')
def vincular_componente_ot(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    c=db.session.get(OTComponente,int(request.form.get('componente_id') or 0)) or abort(404)
    if c.ot_id!=ot.id: abort(400)
    modelo_id=int(request.form.get('modelo_componente_id') or 0)
    modelo=db.session.get(ComponenteModelo,modelo_id) if modelo_id else None
    if modelo_id and (not modelo or not modelo.activo): abort(400)
    if OTReparacionSeleccionada.query.filter_by(ot_componente_id=c.id).first() or PlanActividad.query.filter_by(ot_componente_id=c.id).first():
        flash('Este componente ya tiene reparaciones seleccionadas o planificadas. Revise las asignaciones antes de cambiar el vínculo.')
        return redirect(url_for('ver_ot',ot_id=ot.id))
    c.componente_modelo_id=modelo.id if modelo else None
    db.session.commit(); audit('Vínculo de componente maestro',f'{ot.numero}: {c.nombre}')
    flash('Componente vinculado al maestro. Ahora seleccione sus reparaciones en Diagnóstico.')
    return redirect(url_for('ver_ot',ot_id=ot.id))

@app.route('/ot/<int:ot_id>/diagnostico', methods=['GET','POST'])
@permiso_required('diagnostico')
def diagnostico_ot(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    if request.method=='POST':
        comp_id=request.form.get('ot_componente_id')
        comp=db.session.get(OTComponente,int(comp_id)) if comp_id else None
        if not comp or comp.ot_id!=ot.id or (comp.condicion_ingreso or '').strip().casefold()!='recibido': flash('Solo puede registrar hallazgos de componentes efectivamente recibidos en esta OT.'); return redirect(url_for('diagnostico_ot',ot_id=ot.id))
        condicion=request.form.get('condicion','').strip()
        if not condicion:
            flash('Debes indicar la condición encontrada.'); return redirect(url_for('diagnostico_ot',ot_id=ot.id))
        nombre,mime,datos=preparar_foto(request.files.get('foto_camara') or request.files.get('foto_archivo') or request.files.get('foto'))
        h=DiagnosticoHallazgo(ot_id=ot.id,ot_componente_id=comp.id if comp else None,zona=request.form.get('zona','').strip(),condicion=condicion,trabajo_recomendado=request.form.get('trabajo_recomendado','').strip(),observacion=request.form.get('observacion','').strip(),estado=request.form.get('estado','Pendiente'),foto_nombre=nombre,foto_mime=mime,foto_datos=datos,creado_por=usuario_actual().nombre)
        db.session.add(h); db.session.flush()
        for archivo in request.files.getlist('fotos_adicionales'):
            n,m,d=preparar_foto(archivo)
            if d: db.session.add(DiagnosticoFoto(hallazgo_id=h.id,nombre=n,mime=m,datos=d))
        db.session.commit(); audit('Hallazgo diagnóstico creado',f'{ot.numero} / {comp.nombre if comp else "General"}')
        flash('Hallazgo agregado al diagnóstico.'); return redirect(url_for('diagnostico_ot',ot_id=ot.id))
    marcas_guardadas={m.hallazgo_id:json.loads(m.marcas_json or '[]') for m in DiagnosticoMarca.query.filter(DiagnosticoMarca.hallazgo_id.in_([h.id for h in ot.hallazgos])).all()} if ot.hallazgos else {}
    return render_template('diagnostico.html',ot=ot,componentes_recibidos=[c for c in ot.componentes_detalle if (c.condicion_ingreso or '').strip().casefold()=='recibido'],marcas_guardadas=marcas_guardadas)

@app.route('/diagnostico/hallazgo/<int:hid>/marcas',methods=['POST'])
@permiso_required('diagnostico')
def diagnostico_marcas(hid):
    h=db.session.get(DiagnosticoHallazgo,hid) or abort(404)
    if not h.foto_datos: return jsonify({'error':'Primero debe cargar una fotografía.'}),400
    payload=request.get_json(silent=True) or {}
    marcas=payload.get('marcas',[])
    if not isinstance(marcas,list) or len(marcas)>100: return jsonify({'error':'Marcas inválidas'}),400
    permitidos={'circle','arrow','square'}
    for m in marcas:
        if not isinstance(m,dict) or m.get('tipo') not in permitidos or m.get('color') not in ('#e53935','#e6a700','#1976d2') or not all(isinstance(m.get(k),(int,float)) and 0<=m[k]<=1 for k in ('x','y')):
            return jsonify({'error':'Marcas inválidas'}),400
    registro=DiagnosticoMarca.query.filter_by(hallazgo_id=hid).first()
    if not registro:
        registro=DiagnosticoMarca(hallazgo_id=hid);db.session.add(registro)
    registro.marcas_json=json.dumps(marcas);db.session.commit()
    audit('Anotaciones diagnóstico guardadas',f'{h.ot.numero} / Hallazgo {hid}')
    return jsonify({'ok':True})

@app.route('/diagnostico/hallazgo/<int:hid>/foto')
@login_required
def diagnostico_foto(hid):
    h=db.session.get(DiagnosticoHallazgo,hid) or abort(404)
    if not h.foto_datos: abort(404)
    return send_file(BytesIO(h.foto_datos),mimetype=h.foto_mime or 'image/jpeg',download_name=h.foto_nombre or f'hallazgo-{h.id}.jpg')

@app.route('/diagnostico/foto/<int:foto_id>')
@login_required
def diagnostico_foto_adicional(foto_id):
    foto=db.session.get(DiagnosticoFoto,foto_id) or abort(404)
    return send_file(BytesIO(foto.datos),mimetype=foto.mime or 'image/jpeg',download_name=foto.nombre or f'foto-{foto.id}.jpg')

@app.route('/diagnostico/foto/<int:foto_id>/marcas',methods=['POST'])
@permiso_required('diagnostico')
def diagnostico_foto_marcas(foto_id):
    foto=db.session.get(DiagnosticoFoto,foto_id) or abort(404)
    payload=request.get_json(silent=True) or {}
    marcas=payload.get('marcas',[])
    if not isinstance(marcas,list) or len(marcas)>100: return jsonify({'error':'Marcas inválidas'}),400
    for m in marcas:
        if not isinstance(m,dict) or m.get('tipo') not in ('circle','arrow','square') or m.get('color') not in ('#e53935','#e6a700','#1976d2') or not all(isinstance(m.get(k),(int,float)) and 0<=m[k]<=1 for k in ('x','y')): return jsonify({'error':'Marcas inválidas'}),400
    foto.marcas_json=json.dumps(marcas);db.session.commit()
    return jsonify({'ok':True})

@app.route('/diagnostico/foto/<int:foto_id>/eliminar',methods=['POST'])
@permiso_required('diagnostico')
def diagnostico_foto_eliminar(foto_id):
    foto=db.session.get(DiagnosticoFoto,foto_id) or abort(404)
    ot_id=foto.hallazgo.ot_id
    db.session.delete(foto);db.session.commit()
    flash('Fotografía eliminada.');return redirect(url_for('diagnostico_ot',ot_id=ot_id))

@app.route('/diagnostico/hallazgo/<int:hid>/editar',methods=['GET','POST'])
@permiso_required('diagnostico')
def diagnostico_editar(hid):
    h=db.session.get(DiagnosticoHallazgo,hid) or abort(404)
    ot=h.ot
    if request.method=='POST':
        comp_id=request.form.get('ot_componente_id')
        comp=db.session.get(OTComponente,int(comp_id)) if comp_id else None
        if not comp or comp.ot_id!=ot.id or (comp.condicion_ingreso or '').strip().casefold()!='recibido': abort(400)
        condicion=request.form.get('condicion','').strip()
        if not condicion:
            flash('Debes indicar la condición encontrada.'); return redirect(url_for('diagnostico_editar',hid=h.id))
        h.ot_componente_id=comp.id if comp else None
        h.zona=request.form.get('zona','').strip(); h.condicion=condicion
        h.trabajo_recomendado=request.form.get('trabajo_recomendado','').strip(); h.observacion=request.form.get('observacion','').strip(); h.estado=request.form.get('estado','Pendiente')
        if request.files.get('foto') and request.files['foto'].filename:
            nombre,mime,datos=preparar_foto(request.files.get('foto'))
            if datos: h.foto_nombre=nombre; h.foto_mime=mime; h.foto_datos=datos
        for archivo in request.files.getlist('fotos_adicionales'):
            n,m,d=preparar_foto(archivo)
            if d: db.session.add(DiagnosticoFoto(hallazgo_id=h.id,nombre=n,mime=m,datos=d))
        db.session.commit(); audit('Hallazgo diagnóstico editado',f'{ot.numero} / Hallazgo {h.id}')
        flash('Hallazgo actualizado.'); return redirect(url_for('diagnostico_ot',ot_id=ot.id))
    return render_template('diagnostico_editar.html',ot=ot,h=h)

@app.route('/diagnostico/hallazgo/<int:hid>/eliminar',methods=['POST'])
@permiso_required('diagnostico')
def diagnostico_eliminar(hid):
    h=db.session.get(DiagnosticoHallazgo,hid) or abort(404); ot_id=h.ot_id; detalle=f'{h.ot.numero} / Hallazgo {h.id}'
    DiagnosticoMarca.query.filter_by(hallazgo_id=hid).delete()
    db.session.delete(h); db.session.commit(); audit('Hallazgo diagnóstico eliminado',detalle); flash('Hallazgo eliminado.')
    return redirect(url_for('diagnostico_ot',ot_id=ot_id))

@app.route('/ot/<int:ot_id>/diagnostico/pdf')
@permiso_required('diagnostico')
def diagnostico_pdf(ot_id):
    """V7.40: informe A4 por OT, agrupado por componentes efectivamente recibidos."""
    from html import escape
    from collections import defaultdict
    from reportlab.platypus import KeepTogether, CondPageBreak, HRFlowable
    from reportlab.lib.enums import TA_CENTER

    ot = db.session.get(OT, ot_id) or abort(404)
    azul = colors.HexColor('#183e60')
    texto = colors.HexColor('#2b4153')
    suave = colors.HexColor('#e9eff5')
    borde = colors.HexColor('#dce4eb')
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=34, leftMargin=34,
                            topMargin=49, bottomMargin=44,
                            title=f'Diagnóstico técnico {ot.numero}', author='INGEPRO División Maestranza')
    ancho = A4[0] - 68
    estilos = {
        'normal': ParagraphStyle('v740_normal', fontName='Helvetica', fontSize=8.3, leading=12, textColor=texto, spaceAfter=0),
        'peque': ParagraphStyle('v740_peque', fontName='Helvetica', fontSize=7.4, leading=10.5, textColor=texto),
        'etiqueta': ParagraphStyle('v740_etiqueta', fontName='Helvetica-Bold', fontSize=7.1, leading=10, textColor=azul),
        'titulo': ParagraphStyle('v740_titulo', fontName='Helvetica-Bold', fontSize=12, leading=15, textColor=azul),
        'hallazgo': ParagraphStyle('v740_hallazgo', fontName='Helvetica-Bold', fontSize=9.4, leading=13, textColor=azul),
        'blanco': ParagraphStyle('v740_blanco', fontName='Helvetica-Bold', fontSize=10, leading=13, textColor=colors.white),
        'pie_foto': ParagraphStyle('v740_piefoto', fontName='Helvetica', fontSize=7, leading=9, textColor=colors.HexColor('#65788b'), alignment=TA_CENTER),
    }
    def P(valor, estilo='normal'):
        return Paragraph(escape(str(valor if valor is not None else '-')).replace('\n','<br/>'), estilos[estilo])
    def seccion(nombre):
        return [Spacer(1, 12), P(nombre,'titulo'), Spacer(1, 7)]
    def tabla_estilo(t, fondo=None, pad=8):
        reglas=[('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),pad),
                ('RIGHTPADDING',(0,0),(-1,-1),pad),('TOPPADDING',(0,0),(-1,-1),pad),
                ('BOTTOMPADDING',(0,0),(-1,-1),pad)]
        if fondo: reglas.append(('BACKGROUND',(0,0),(-1,-1),fondo))
        t.setStyle(TableStyle(reglas))
        return t
    def foto_marcada(datos, marcas):
        try:
            im=PILImage.open(BytesIO(datos))
            im=ImageOps.exif_transpose(im).convert('RGB')
            if marcas:
                dibujo=ImageDraw.Draw(im)
                w,h=im.size
                for marca in marcas:
                    try:
                        x=int(float(marca.get('x',0))*w); y=int(float(marca.get('y',0))*h)
                        r=max(12,int(w*.055)); grosor=max(3,int(w*.006))
                        color=marca.get('color','#e03131')
                        if not re.fullmatch(r'#[0-9a-fA-F]{6}',str(color)): color='#e03131'
                        tipo=marca.get('tipo')
                        if tipo=='circle': dibujo.ellipse((x-r,y-r,x+r,y+r),outline=color,width=grosor)
                        elif tipo=='square': dibujo.rectangle((x-r,y-r,x+r,y+r),outline=color,width=grosor)
                        elif tipo=='arrow':
                            dibujo.line([(x-r,y+r),(x+r,y-r)],fill=color,width=grosor)
                            dibujo.line([(x+r,y-r),(x,y-r//2)],fill=color,width=grosor)
                            dibujo.line([(x+r,y-r),(x+r//2,y)],fill=color,width=grosor)
                    except (TypeError,ValueError,KeyError): continue
            # Redondeo sutil aplicado sobre los píxeles de la fotografía.
            # Se conserva toda la imagen y su proporción; solo se suavizan las cuatro esquinas.
            radio=max(2, round(min(im.size)*0.012))
            mascara=PILImage.new('L',im.size,0)
            ImageDraw.Draw(mascara).rounded_rectangle((0,0,im.width-1,im.height-1),radius=radio,fill=255)
            rgba=im.convert('RGBA'); rgba.putalpha(mascara)
            salida=BytesIO(); rgba.save(salida,format='PNG',optimize=True)
            return salida.getvalue(),im.size
        except Exception: return None,None
    def imagen_con_pie(datos,marcas,numero):
        datos_procesados,tamano=foto_marcada(datos,marcas)
        if not datos_procesados: return None
        w,h=tamano
        # La fotografía ocupa como máximo el ancho útil de su columna,
        # sin recortar, deformar ni alterar su relación ancho/alto.
        columna_foto=ancho*.45
        limite_ancho=columna_foto-14
        limite_alto=210
        escala=min(limite_ancho/w,limite_alto/h)
        img=Image(BytesIO(datos_procesados),width=w*escala,height=h*escala)
        img.hAlign='CENTER'
        pie=P(f'Figura {numero:02d} · Registro fotográfico','pie_foto')
        # Celda interna de ancho completo: centra fotos verticales y horizontales.
        centrado=Table([[img],[pie]],colWidths=[columna_foto-7])
        centrado.setStyle(TableStyle([
            ('ALIGN',(0,0),(-1,-1),'CENTER'),
            ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),
            ('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(0,0),5),
            ('BOTTOMPADDING',(0,1),(0,1),0)
        ]))
        return centrado
    def pie_pagina(canvas, documento):
        canvas.saveState(); w,h=A4
        canvas.setStrokeColor(borde);canvas.line(34,35,w-34,35)
        canvas.setFont('Helvetica',7);canvas.setFillColor(colors.HexColor('#738395'))
        canvas.drawString(34,24,'INGEPRO · DIVISIÓN MAESTRANZA · Documento técnico')
        canvas.drawRightString(w-34,24,f'OT {ot.numero} · Página {documento.page}')
        if documento.page>1:
            canvas.setFont('Helvetica-Bold',8);canvas.setFillColor(azul)
            canvas.drawString(34,h-29,'INGEPRO  |  INFORME DE DIAGNÓSTICO TÉCNICO')
            canvas.setStrokeColor(azul);canvas.line(34,h-35,w-34,h-35)
        canvas.restoreState()

    story=[]
    logo_path=os.path.join(app.root_path,'static','img','logo-ingepro.png')
    if os.path.exists(logo_path):
        try:
            with PILImage.open(logo_path) as im: lw,lh=im.size
            factor=min(138/lw,53/lh)
            logo=Image(logo_path,width=lw*factor,height=lh*factor)
        except Exception: logo=P('INGEPRO','titulo')
    else: logo=P('INGEPRO','titulo')
    cabecera=Table([[logo,Paragraph('<b>INFORME DE DIAGNÓSTICO TÉCNICO</b><br/>'+
                     escape(ot.numero)+' · Rev. 00<br/>'+datetime.now().strftime('%d/%m/%Y'),
                     ParagraphStyle('v740_head',fontName='Helvetica',fontSize=9,leading=14,textColor=azul,alignment=TA_RIGHT))]],
                   colWidths=[ancho*.33,ancho*.67])
    cabecera.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('BOTTOMPADDING',(0,0),(-1,-1),12),
                                  ('LINEBELOW',(0,0),(-1,-1),3,azul)]))
    story.extend([cabecera,Spacer(1,13)])
    recibidos=[c for c in ot.componentes_detalle if (c.condicion_ingreso or '').strip().casefold()=='recibido']
    ids_recibidos={c.id for c in recibidos}
    hallazgos=[h for h in ot.hallazgos if h.ot_componente_id in ids_recibidos]
    agrupados=defaultdict(list)
    for h in hallazgos: agrupados[h.ot_componente_id].append(h)
    story.extend(seccion('01. Antecedentes generales'))
    campos=[('Cliente',ot.cliente),('OT INGEPRO',ot.numero),('Equipo',ot.equipo),
            ('Serie / Tag',ot.serie or '-'),('Guía de ingreso',ot.guia or '-'),
            ('Componentes recibidos',str(len(recibidos))),('Hallazgos registrados',str(len(hallazgos))),
            ('Fecha de emisión',datetime.now().strftime('%d/%m/%Y'))]
    celdas=[]
    for i in range(0,len(campos),2):
        fila=[]
        for etiqueta,valor in campos[i:i+2]:
            fila.append([P(etiqueta,'etiqueta'),Spacer(1,3),P(valor,'normal')])
        celdas.append(fila)
    meta=Table(celdas,colWidths=[ancho/2]*2,cornerRadii=[9,9,9,9]);tabla_estilo(meta,colors.HexColor('#f0f4f8'),9)
    story.append(meta)
    story.extend(seccion('02. Objetivo del diagnóstico'))
    story.append(P('Documentar las condiciones identificadas durante la inspección de los componentes efectivamente recibidos en maestranza, con evidencia fotográfica, observaciones técnicas y recomendaciones de reparación.'))
    if ot.diagnostico:
        story.extend(seccion('Resumen general registrado en OT'))
        story.append(P(ot.diagnostico))
    story.extend(seccion('03. Diagnóstico por componente'))
    if not recibidos: story.append(P('No existen componentes registrados como recibidos en esta OT.'))
    contador_fotos=0
    for indice,comp in enumerate(recibidos,1):
        lista=agrupados.get(comp.id,[])
        encabezado=Table([[P(f'COMPONENTE {indice:02d} · {comp.nombre.upper()}','blanco')],
                           [Paragraph(escape(f'Código: {comp.codigo or "-"}  ·  {len(lista)} hallazgo(s)'),
                            ParagraphStyle(f'v740_comp_{indice}',fontName='Helvetica',fontSize=7.5,leading=11,textColor=colors.HexColor('#d6e4ee')))]],
                         colWidths=[ancho],cornerRadii=[7,7,7,7])
        tabla_estilo(encabezado,azul,7)
        story.extend([CondPageBreak(100),Spacer(1,10),encabezado,Spacer(1,8)])
        if not lista:
            story.append(P('Componente recibido sin hallazgos registrados.'))
            continue
        for numero,hallazgo in enumerate(lista,1):
            fotos=[]
            if hallazgo.foto_datos:
                registro=DiagnosticoMarca.query.filter_by(hallazgo_id=hallazgo.id).first()
                try: marcas=json.loads(registro.marcas_json or '[]') if registro else []
                except (ValueError,TypeError): marcas=[]
                fotos.append((hallazgo.foto_datos,marcas))
            for foto in hallazgo.fotos_adicionales:
                try: marcas=json.loads(foto.marcas_json or '[]')
                except (ValueError,TypeError): marcas=[]
                fotos.append((foto.datos,marcas))
            comentario=[P('CONDICIÓN ENCONTRADA','etiqueta'),Spacer(1,3),P(hallazgo.condicion),Spacer(1,8),
                        HRFlowable(width='100%',thickness=.5,color=colors.HexColor('#cad8e4')),Spacer(1,7),
                        P('TRABAJO RECOMENDADO','etiqueta'),Spacer(1,3),P(hallazgo.trabajo_recomendado or '-'),Spacer(1,8),
                        HRFlowable(width='100%',thickness=.5,color=colors.HexColor('#cad8e4')),Spacer(1,7),
                        P('COMENTARIO TÉCNICO','etiqueta'),Spacer(1,3),P(hallazgo.observacion or '-')]
            panel=Table([[comentario]],colWidths=[ancho*.55-16],cornerRadii=[9,9,9,9]);tabla_estilo(panel,suave,9)
            titulo=P(f'Hallazgo {numero:02d} · {hallazgo.zona or "Inspección técnica"}','hallazgo')
            if not fotos:
                izquierda=[P('Sin fotografía registrada','peque')]
                fila=Table([[izquierda,panel]],colWidths=[ancho*.45,ancho*.55])
                fila.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),7)]))
                story.append(KeepTogether([titulo,Spacer(1,6),fila,Spacer(1,12)]))
            else:
                for foto_indice,(datos,marcas) in enumerate(fotos):
                    contador_fotos+=1
                    izquierda=imagen_con_pie(datos,marcas,contador_fotos) or P('Fotografía no disponible','peque')
                    derecha=panel if foto_indice==0 else P(f'Continuación del hallazgo {numero:02d} · Evidencia fotográfica adicional','peque')
                    fila=Table([[izquierda,derecha]],colWidths=[ancho*.45,ancho*.55])
                    fila.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),
                                              ('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),4),
                                              ('BOTTOMPADDING',(0,0),(-1,-1),4)]))
                    bloque=[titulo if foto_indice==0 else P(f'Hallazgo {numero:02d} · Fotografía adicional','hallazgo'),Spacer(1,6),fila,Spacer(1,10)]
                    # Mantener fotografía y comentario juntos en la misma página.
                    story.append(KeepTogether(bloque))
    story.extend(seccion('04. Resumen consolidado del diagnóstico'))
    resumen=[[P('COMPONENTE','etiqueta'),P('HALLAZGOS','etiqueta'),P('ESTADO','etiqueta')]]
    for c in recibidos:
        hs=agrupados.get(c.id,[])
        resumen.append([P(c.nombre,'peque'),P(str(len(hs)),'peque'),P(', '.join(sorted({h.estado or 'Pendiente' for h in hs})) if hs else 'Sin hallazgos','peque')])
    t=Table(resumen,colWidths=[ancho*.48,ancho*.18,ancho*.34],repeatRows=1,cornerRadii=[7,7,7,7]);tabla_estilo(t,None,7)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),suave),('LINEBELOW',(0,0),(-1,-1),.5,borde)]));story.append(t)
    story.extend(seccion('05. Conclusión general de la OT'))
    # La conclusión se toma de la OT, no se inventa una evaluación automática.
    conclusion=ot.diagnostico or 'Conclusión general pendiente de registro por el área técnica.'
    t=Table([[[P('EVALUACIÓN TÉCNICA CONSOLIDADA','etiqueta'),Spacer(1,5),P(conclusion)]]],
            colWidths=[ancho],cornerRadii=[9,9,9,9]);tabla_estilo(t,suave,10);story.append(t)
    story.extend(seccion('06. Responsables del informe'))
    firmas=Table([[[P('ELABORADO POR','etiqueta'),Spacer(1,28),HRFlowable(width='95%',thickness=.6,color=borde),
                     Spacer(1,5),P('Técnico responsable · Nombre / Firma / Fecha','peque')],
                   [P('REVISADO POR','etiqueta'),Spacer(1,28),HRFlowable(width='95%',thickness=.6,color=borde),
                     Spacer(1,5),P('Área de Ingeniería · Nombre / Firma / Fecha','peque')]]],
                 colWidths=[ancho/2]*2,cornerRadii=[9,9,9,9]);tabla_estilo(firmas,colors.HexColor('#f0f4f8'),10)
    story.append(firmas)
    story.extend([Spacer(1,12),P('Este informe documenta las condiciones observadas durante el diagnóstico. Las reparaciones están sujetas a evaluación técnica, aprobación del alcance y condiciones contractuales aplicables.','peque')])
    doc.build(story,onFirstPage=pie_pagina,onLaterPages=pie_pagina)
    buf.seek(0)
    audit('PDF diagnóstico generado',ot.numero)
    return send_file(buf,mimetype='application/pdf',as_attachment=True,download_name=f'Diagnostico-{ot.numero}.pdf')

@app.route('/clientes')
@permiso_required('clientes')
def clientes(): return render_template('clientes.html',clientes=Cliente.query.order_by(Cliente.nombre).all())

@app.route('/clientes/nuevo',methods=['GET','POST'])
@permiso_required('clientes')
def cliente_nuevo():
    if request.method=='POST':
        rut=normalizar_rut(request.form.get('rut'))
        if not rut_valido(rut): flash('RUT inválido. Revisa el dígito verificador.'); return render_template('cliente_form.html',cliente=None)
        existente=Cliente.query.filter_by(rut=rut).first()
        if existente: flash(f'El RUT {formatear_rut(rut)} ya está registrado en {existente.nombre}.'); return redirect(url_for('cliente_ver',cid=existente.id))
        c=Cliente(nombre=request.form['nombre'].strip(),nombre_comercial=request.form.get('nombre_comercial','').strip(),rut=rut,direccion=request.form.get('direccion','').strip(),comuna=request.form.get('comuna','').strip(),ciudad=request.form.get('ciudad','').strip(),giro=request.form.get('giro','').strip(),telefono=request.form.get('telefono','').strip(),email=request.form.get('email','').strip())
        db.session.add(c); db.session.commit(); audit('Cliente creado',f'{c.nombre} / {c.rut}'); return redirect(url_for('cliente_ver',cid=c.id))
    return render_template('cliente_form.html',cliente=None)

@app.route('/clientes/<int:cid>',methods=['GET','POST'])
@permiso_required('clientes')
def cliente_ver(cid):
    c=db.session.get(Cliente,cid) or abort(404)
    if request.method=='POST':
        if request.form.get('accion')=='contacto':
            x=ContactoCliente(cliente_id=c.id,nombre=request.form['contacto_nombre'].strip(),cargo=request.form.get('cargo','').strip(),email=request.form.get('contacto_email','').strip(),telefono=request.form.get('contacto_telefono','').strip(),recibe_cotizacion=bool(request.form.get('recibe_cotizacion')),puede_aprobar=bool(request.form.get('puede_aprobar')),recibe_diagnostico=bool(request.form.get('recibe_diagnostico')),recibe_informe_final=bool(request.form.get('recibe_informe_final')))
            db.session.add(x); db.session.commit(); audit('Contacto cliente creado',f'{c.nombre}: {x.nombre}'); flash('Contacto agregado.')
        return redirect(url_for('cliente_ver',cid=c.id))
    return render_template('cliente_ver.html',cliente=c,formatear_rut=formatear_rut)

def prefijo_codigo(valor, largo):
    limpio=unicodedata.normalize('NFKD', (valor or '').upper()).encode('ascii','ignore').decode()
    letras=re.sub(r'[^A-Z]', '', limpio)
    return (letras + 'X'*largo)[:largo]

def codigo_equipo(m):
    return f'{prefijo_codigo(m.tipo,2)}-{prefijo_codigo(m.marca,3)}-{m.id:04d}'

def codigo_componente(c, m):
    return f'{prefijo_codigo(c.nombre,2)}-{prefijo_codigo(m.marca,3)}-{c.id:04d}'

@app.route('/maestros/equipos',methods=['GET','POST'])
@permiso_required('maestros')
def maestros_equipos():
    if request.method=='POST':
        tipo=request.form.get('tipo','').strip(); marca=request.form.get('marca','').strip(); modelo=request.form.get('modelo','').strip()
        if not all((tipo,marca,modelo)):
            flash('Complete tipo, marca y modelo.'); return redirect(url_for('maestros_equipos'))
        existe=next((x for x in ModeloEquipo.query.all() if (x.tipo.casefold(),x.marca.casefold(),x.modelo.casefold())==(tipo.casefold(),marca.casefold(),modelo.casefold())),None)
        if existe:
            flash('Ya existe ese tipo, marca y modelo.'); return redirect(url_for('maestro_modelo',mid=existe.id))
        m=ModeloEquipo(tipo=tipo,marca=marca,modelo=modelo)
        db.session.add(m); db.session.commit(); audit('Modelo de equipo creado',f'{codigo_equipo(m)}: {m.marca} {m.modelo}')
        return redirect(url_for('maestro_modelo',mid=m.id))
    return render_template('maestros_equipos.html',modelos=ModeloEquipo.query.filter_by(activo=True).order_by(ModeloEquipo.marca,ModeloEquipo.modelo).all(),codigo_equipo=codigo_equipo)

@app.route('/maestros/equipos/<int:mid>',methods=['GET','POST'])
@permiso_required('maestros')
def maestro_modelo(mid):
    m=db.session.get(ModeloEquipo,mid) or abort(404)
    if request.method=='POST':
        nombre=request.form.get('nombre','').strip()
        if not nombre: abort(400)
        c=ComponenteModelo(modelo_id=m.id,nombre=nombre,categoria=request.form.get('categoria','Componente reparable'),material_default=request.form.get('material_default','').strip())
        db.session.add(c); db.session.flush(); c.codigo=codigo_componente(c,m)
        db.session.commit(); audit('Componente de modelo creado',f'{m.marca} {m.modelo}: {c.codigo} {c.nombre}')
        flash('Componente agregado al despiece.'); return redirect(url_for('maestro_modelo',mid=m.id))
    return render_template('maestro_modelo.html',modelo=m,codigo_equipo=codigo_equipo)

@app.route('/maestros/equipos/<int:mid>/editar',methods=['POST'])
@permiso_required('maestros')
def editar_modelo(mid):
    m=db.session.get(ModeloEquipo,mid) or abort(404)
    nuevo=request.form.get('modelo','').strip()
    if not nuevo: abort(400)
    if any(x.id!=m.id and (x.tipo.casefold(),x.marca.casefold(),x.modelo.casefold())==(m.tipo.casefold(),m.marca.casefold(),nuevo.casefold()) for x in ModeloEquipo.query.all()):
        flash('El modelo ya existe.'); return redirect(url_for('maestro_modelo',mid=mid))
    anterior=m.modelo; m.modelo=nuevo; db.session.commit(); audit('Modelo editado',f'{anterior} → {nuevo}')
    return redirect(url_for('maestro_modelo',mid=mid))

@app.route('/maestros/componentes/<int:cid>/editar',methods=['POST'])
@permiso_required('maestros')
def editar_componente(cid):
    c=db.session.get(ComponenteModelo,cid) or abort(404)
    nombre=request.form.get('nombre','').strip()
    if not nombre: abort(400)
    c.nombre=nombre; c.material_default=request.form.get('material_default','').strip()
    c.categoria=request.form.get('categoria',c.categoria)
    db.session.commit(); audit('Componente editado',f'{c.id}: {nombre}')
    return redirect(url_for('maestro_modelo',mid=c.modelo_id))

@app.route('/maestros/equipos/<int:mid>/eliminar',methods=['POST'])
@permiso_required('maestros')
def eliminar_modelo(mid):
    if not usuario_actual().es_admin: abort(403)
    m=db.session.get(ModeloEquipo,mid) or abort(404)
    if request.form.get('confirmacion')!='ELIMINAR':
        flash('Escriba ELIMINAR para confirmar.'); return redirect(url_for('maestro_modelo',mid=mid))
    if OT.query.filter_by(modelo_id=mid).first() or OTComponente.query.filter(OTComponente.componente_modelo_id.in_([c.id for c in m.componentes])).first():
        m.activo=False; db.session.commit(); audit('Modelo desactivado por referencias históricas',codigo_equipo(m)); flash('Modelo desactivado: existen referencias de OT.')
    else:
        codigo=codigo_equipo(m); db.session.delete(m); db.session.commit(); audit('Modelo eliminado',codigo); flash('Modelo eliminado.')
    return redirect(url_for('maestros_equipos'))

@app.route('/maestros/componentes/<int:cid>/eliminar',methods=['POST'])
@permiso_required('maestros')
def eliminar_componente(cid):
    if not usuario_actual().es_admin: abort(403)
    c=db.session.get(ComponenteModelo,cid) or abort(404)
    if request.form.get('confirmacion')!='ELIMINAR': abort(400)
    mid=c.modelo_id
    if OTComponente.query.filter_by(componente_modelo_id=cid).first() or OTReparacionSeleccionada.query.filter(OTReparacionSeleccionada.reparacion_maestra_id.in_([r.id for r in c.reparaciones])).first():
        c.activo=False; db.session.commit(); audit('Componente desactivado',str(cid)); flash('Componente desactivado: tiene referencias históricas.')
    else:
        db.session.delete(c); db.session.commit(); audit('Componente eliminado',str(cid)); flash('Componente eliminado.')
    return redirect(url_for('maestro_modelo',mid=mid))

@app.route('/maestros/componentes/<int:cid>/reparaciones',methods=['GET','POST'])
@permiso_required('maestros')
def maestro_reparaciones(cid):
    c=db.session.get(ComponenteModelo,cid) or abort(404)
    if request.method=='POST':
        r=ReparacionMaestra(componente_modelo_id=c.id,nombre=request.form['nombre'].strip(),descripcion=request.form.get('descripcion','').strip(),proceso=request.form.get('proceso','').strip(),hh_estimadas=float(request.form.get('hh_estimadas') or 0),recurso=request.form.get('recurso','').strip(),maquina=request.form.get('maquina','').strip(),secuencia=int(request.form.get('secuencia') or 0))
        db.session.add(r); db.session.commit(); audit('Reparación maestra creada',f'{c.nombre}: {r.nombre}'); flash('Reparación agregada al componente.'); return redirect(url_for('maestro_reparaciones',cid=c.id))
    return render_template('maestro_reparaciones.html',componente=c,reparaciones=ReparacionMaestra.query.filter_by(componente_modelo_id=c.id).order_by(ReparacionMaestra.secuencia,ReparacionMaestra.nombre).all())

CPR_112A20_TAREAS = [
 'ARENADO Y PINTURA ANTICORROSIVA','RETIRO DE PERNOS CORTADOS',
 'MECANIZADO ALOJAMIENTO RETEN','FABRICACION Y MONTAJE CAMISA ALOJAMIENTO RETEN',
 'PREMECANIZADO CAMISAS','MECANIZADO ALOJAMIENTO 02',
 'FABRICACION E INSTALACIÓN DE CAMISAS','REPARAR HILOS CON INSERTO',
 'CONTROL DE CALIDAD ALOJAMIENTO Y RETEN','LIMPIEZA PINTURA',
 'REPASO DE HILOS','CHECK LIST']

@app.route('/maestros/componentes/<int:cid>/cargar-cpr-112a20',methods=['POST'])
@permiso_required('maestros')
def cargar_cpr_112a20(cid):
    c=db.session.get(ComponenteModelo,cid) or abort(404)
    if 'CPR' not in c.nombre.upper() and 'RODAMIENTO' not in c.nombre.upper():
        flash('Esta plantilla solo corresponde a componentes CPR / cuerpo de rodamientos.')
        return redirect(url_for('maestro_reparaciones',cid=cid))
    existentes={r.nombre.strip().upper() for r in ReparacionMaestra.query.filter_by(componente_modelo_id=cid).all()}
    nuevos=0
    for n,nombre in enumerate(CPR_112A20_TAREAS,1):
        if nombre.upper() not in existentes:
            db.session.add(ReparacionMaestra(componente_modelo_id=cid,nombre=nombre,secuencia=n,activo=True))
            nuevos+=1
    db.session.commit()
    flash(f'Plantilla CPR 112A20: {nuevos} tareas nuevas. Revise horas y recursos antes de cotizar.')
    return redirect(url_for('maestro_reparaciones',cid=cid))

@app.route('/ot/<int:ot_id>/reparaciones',methods=['GET','POST'])
@permiso_required('diagnostico')
def ot_reparaciones(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    componentes=[c for c in ot.componentes_detalle if (c.condicion_ingreso or '').strip().lower()=='recibido']
    disponibles={c.id:ReparacionMaestra.query.filter_by(componente_modelo_id=c.componente_modelo_id,activo=True).order_by(ReparacionMaestra.secuencia,ReparacionMaestra.id).all() if c.componente_modelo_id else [] for c in componentes}
    if request.method=='POST':
        permitidos={(c.id,r.id) for c in componentes for r in disponibles[c.id]}
        marcados=set()
        for valor in request.form.getlist('reparacion'):
            try: cid,rid=map(int,valor.split(':'))
            except (ValueError,TypeError): abort(400)
            if (cid,rid) not in permitidos: abort(400)
            marcados.add((cid,rid))
        existentes=OTReparacionSeleccionada.query.filter_by(ot_id=ot.id).all()
        for x in existentes:
            if (x.ot_componente_id,x.reparacion_maestra_id) not in marcados:
                if PlanActividad.query.filter_by(ot_id=ot.id,ot_componente_id=x.ot_componente_id,reparacion_maestra_id=x.reparacion_maestra_id).first():
                    flash('Una tarea ya planificada no puede desmarcarse; elimine o reprograme primero su actividad.')
                    return redirect(url_for('ot_reparaciones',ot_id=ot.id))
                db.session.delete(x)
        prev={(x.ot_componente_id,x.reparacion_maestra_id) for x in existentes}
        for cid,rid in marcados-prev:
            db.session.add(OTReparacionSeleccionada(ot_id=ot.id,ot_componente_id=cid,reparacion_maestra_id=rid))
        db.session.commit(); flash('Reparaciones necesarias guardadas para esta OT.')
        return redirect(url_for('ot_reparaciones',ot_id=ot.id))
    seleccionadas={(x.ot_componente_id,x.reparacion_maestra_id) for x in OTReparacionSeleccionada.query.filter_by(ot_id=ot.id).all()}
    return render_template('ot_reparaciones.html',ot=ot,componentes=componentes,disponibles=disponibles,seleccionadas=seleccionadas)

@app.route('/maestros/reparaciones/<int:rid>/estado',methods=['POST'])
@permiso_required('maestros')
def reparacion_estado(rid):
    r=db.session.get(ReparacionMaestra,rid) or abort(404); r.activo=not r.activo; db.session.commit(); return redirect(url_for('maestro_reparaciones',cid=r.componente_modelo_id))

def operarios_planificables():
    return Usuario.query.filter_by(activo=True, disponible_planificacion=True).order_by(Usuario.nombre).all()

def horas_ocupadas(usuario_id, dia, excluir_id=None):
    q=PlanActividad.query.filter(PlanActividad.usuario_id==usuario_id, PlanActividad.fecha_inicio<=dia, PlanActividad.fecha_fin>=dia)
    if excluir_id: q=q.filter(PlanActividad.id!=excluir_id)
    return sum(a.horas_dia or 0 for a in q.all() if a.estado!='Terminada')

def limite_horas_dia(dia):
    """Jornada ordinaria: lunes y martes 9h; miércoles a viernes 8h.
    Los límites provisionales de sábado y domingo se mantienen sin cambios
    hasta acordar una política de fin de semana y horas extraordinarias.
    """
    return 9.0 if dia.weekday() in (0, 1) else 8.0


app.jinja_env.globals['limite_horas_dia'] = limite_horas_dia
app.jinja_env.globals['fecha_dia'] = date

def validar_cupo(usuario, fi, ff, horas, excluir_id=None):
    if not usuario or not usuario.activo or not usuario.disponible_planificacion: return 'Seleccione un trabajador operativo habilitado.'
    if ff<fi or horas<=0: return 'Verifique las fechas y horas de la actividad.'
    dia=fi
    while dia<=ff:
        libres=max(0,limite_horas_dia(dia)-horas_ocupadas(usuario.id,dia,excluir_id))
        if horas>libres+0.0001:
            return f'{usuario.nombre}: el {dia.strftime("%d/%m/%Y")} solo tiene {libres:g} horas disponibles.'
        dia+=timedelta(days=1)
    return None

@app.route('/api/planificacion/disponibilidad')
@permiso_required('planificacion')
def api_plan_disponibilidad():
    try:
        uid=int(request.args.get('usuario_id') or 0); dia=date.fromisoformat(request.args['fecha']); excluir=int(request.args.get('excluir_id') or 0)
    except (ValueError,KeyError): return jsonify({'error':'Parámetros inválidos'}),400
    usuario=db.session.get(Usuario,uid)
    if not usuario or not usuario.activo or not usuario.disponible_planificacion: return jsonify({'error':'Trabajador no habilitado'}),400
    ocupadas=horas_ocupadas(uid,dia,excluir or None)
    return jsonify({'ocupadas':ocupadas,'limite':limite_horas_dia(dia),'disponibles':max(0,limite_horas_dia(dia)-ocupadas),'fecha':dia.isoformat()})

def asignaciones_reparacion(ot_id, componente_id, reparacion):
    """Incluye actividades antiguas sin FK al maestro, por nombre exacto normalizado."""
    import unicodedata
    def normalizar(valor):
        texto=unicodedata.normalize('NFKD', (valor or '').strip().casefold())
        return ' '.join(''.join(c for c in texto if not unicodedata.combining(c)).split())
    nombre=normalizar(reparacion.nombre)
    registros=PlanActividad.query.filter_by(ot_id=ot_id,ot_componente_id=componente_id).all()
    return [a for a in registros if a.reparacion_maestra_id==reparacion.id or
            (a.reparacion_maestra_id is None and normalizar(a.tarea)==nombre)]

@app.route('/gerencia')
@login_required
def gerencia():
    # Información sensible comercial: únicamente administradores o usuarios con permiso de cotizaciones.
    if not (usuario_actual().es_admin or usuario_actual().puede('cotizaciones')): abort(403)
    ots=OT.query.order_by(OT.fecha.desc(), OT.id.desc()).all()
    cotizaciones=Cotizacion.query.filter(Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).all()
    cot_por_ot={}
    for q in cotizaciones: cot_por_ot.setdefault(q.ot_id,q)
    actividades=PlanActividad.query.order_by(PlanActividad.fecha_inicio,PlanActividad.id).all()
    act_por_ot={}
    for a in actividades: act_por_ot.setdefault(a.ot_id,[]).append(a)
    no_aplica=PlanNoAplica.query.all()
    no_aplica_por_comp={}
    for n in no_aplica: no_aplica_por_comp.setdefault(n.ot_componente_id,set()).add(n.reparacion_maestra_id)
    filas=[]
    for ot in ots:
        q=cot_por_ot.get(ot.id)
        actividades_ot=act_por_ot.get(ot.id,[])
        componentes=[]
        for comp in ot.componentes_detalle:
            if (comp.condicion_ingreso or '').strip().casefold()=='no recibido': continue
            tareas=[a for a in actividades_ot if a.ot_componente_id==comp.id]
            etapas={}
            for a in tareas:
                key=a.reparacion_maestra_id or ('manual', (a.tarea or '').strip().casefold())
                etapa=etapas.setdefault(key,{'nombre':a.reparacion.nombre if a.reparacion else (a.tarea or 'Actividad sin descripción'),'asignaciones':[],'estados':set()})
                etapa['asignaciones'].append(a)
                etapa['estados'].add((a.estado or 'Planificada').strip().casefold())
            maestras=ReparacionMaestra.query.filter_by(componente_modelo_id=comp.componente_modelo_id,activo=True).order_by(ReparacionMaestra.secuencia,ReparacionMaestra.id).all() if comp.componente_modelo_id else []
            ignoradas=no_aplica_por_comp.get(comp.id,set())
            for m in maestras:
                if m.id in ignoradas: continue
                etapas.setdefault(m.id,{'nombre':m.nombre,'asignaciones':[],'estados':set()})
            finalizadas=sum(1 for e in etapas.values() if e['asignaciones'] and e['estados'] and e['estados'].issubset({'finalizada','finalizado','terminada','terminado','completada','completado'}))
            activas=sum(1 for e in etapas.values() if e['asignaciones'])
            if not etapas: estado='Sin programación'
            elif finalizadas==len(etapas): estado='Finalizada'
            elif any(('ejecuci' in estado or estado=='en proceso') for e in etapas.values() for estado in e['estados']): estado='En ejecución'
            elif activas: estado='Programada'
            else: estado='Sin programación'
            componentes.append({'componente':comp,'etapas':list(etapas.values()),'estado':estado,'avance':round(finalizadas/len(etapas)*100) if etapas else 0,'programadas':activas,'total':len(etapas)})
        if not componentes and actividades_ot:
            # Actividades antiguas sin componente vinculado: no ocultarlas.
            etapas=[{'nombre':a.reparacion.nombre if a.reparacion else a.tarea,'asignaciones':[a]} for a in actividades_ot if not a.ot_componente_id]
            if etapas: componentes.append({'componente':None,'etapas':etapas,'estado':'Programada','avance':0,'programadas':len(etapas),'total':len(etapas)})
        total=sum(c['total'] for c in componentes)
        final=sum(round(c['avance']*c['total']/100) for c in componentes)
        etapa_actual=next((e['nombre'] for c in componentes for e in c['etapas'] if e['asignaciones'] and not all((a.estado or '').strip().casefold() in ('finalizada','finalizado','terminada','terminado','completada','completado') for a in e['asignaciones'])), None)
        if not etapa_actual: etapa_actual='Sin programación' if not actividades_ot else ('Finalizada' if total and final==total else 'Ver componentes')
        filas.append({'ot':ot,'cot':q,'componentes':componentes,'etapa_actual':etapa_actual,'avance':round(final/total*100) if total else 0,'plan_estado':('Finalizada' if total and final==total else 'En ejecución' if any(c['estado']=='En ejecución' for c in componentes) else 'Programada' if actividades_ot else 'Sin programación')})
    return render_template('gerencia.html',filas=filas)

@app.route('/planificacion',methods=['GET','POST'])
@app.route('/planificacion/detalle',methods=['GET','POST'])
@app.route('/planificacion/administracion',methods=['GET','POST'])
@permiso_required('planificacion')
def planificacion():
    if request.method=='GET' and request.path.rstrip('/') in ('/planificacion', '/planificacion/detalle'):
        return redirect(url_for('planificacion_gantt', **{k:request.args[k] for k in ('mes','anio') if k in request.args}))
    hoy=date.today(); mes=int(request.args.get('mes') or hoy.month); anio=int(request.args.get('anio') or hoy.year)
    if request.method=='POST':
        ot=db.session.get(OT,int(request.form['ot_id'])) or abort(404)
        q_activa=Cotizacion.query.filter(Cotizacion.ot_id==ot.id,Cotizacion.estado=='Aprobada').first()
        if not q_activa:
            flash('La OT requiere una cotización aprobada para planificar trabajos.')
            return redirect(url_for('ver_ot',ot_id=ot.id))
        comp_id=int(request.form.get('ot_componente_id') or 0) or None; rep_id=int(request.form.get('reparacion_maestra_id') or 0) or None
        rep=db.session.get(ReparacionMaestra,rep_id) if rep_id else None
        tarea=(rep.nombre if rep else '').strip(); fi=datetime.strptime(request.form['fecha_inicio'],'%Y-%m-%d').date(); ff=datetime.strptime(request.form['fecha_fin'],'%Y-%m-%d').date()
        try: horas=float(request.form.get('horas_dia') or 0)
        except ValueError: horas=0
        usuario=db.session.get(Usuario,int(request.form.get('usuario_id') or 0))
        error=validar_cupo(usuario,fi,ff,horas)
        if error: flash(error); return redirect(url_for('planificacion',mes=fi.month,anio=fi.year))
        if not comp_id or not OTComponente.query.filter_by(id=comp_id,ot_id=ot.id,condicion_ingreso='Recibido').first():
            flash('Solo puede planificar componentes marcados como Recibido en el despiece de la OT.')
            return redirect(url_for('planificacion',mes=fi.month,anio=fi.year))
        comp=db.session.get(OTComponente,comp_id)
        if not rep or not comp or rep.componente_modelo_id!=comp.componente_modelo_id or not rep.activo:
            flash('Seleccione una reparación activa del maestro correspondiente al componente.')
            return redirect(url_for('planificacion',mes=fi.month,anio=fi.year))
        if PlanNoAplica.query.filter_by(ot_id=ot.id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id).first():
            flash('Esta reparación está marcada como No aplica. Revierta ese estado antes de programar.')
            return redirect(url_for('planificacion',mes=fi.month,anio=fi.year))
        seleccion=OTReparacionSeleccionada.query.filter_by(ot_id=ot.id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id).first()
        if not seleccion:
            db.session.add(OTReparacionSeleccionada(ot_id=ot.id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id))
        # HH del maestro: referencia técnica, nunca límite de programación.
        # La capacidad real del trabajador ya fue validada con validar_cupo().
        if not tarea: flash('Indique una tarea.'); return redirect(url_for('planificacion',mes=fi.month,anio=fi.year))
        a=PlanActividad(ot_id=ot.id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id,usuario_id=int(request.form['usuario_id']),tarea=tarea,fecha_inicio=fi,fecha_fin=ff,horas_dia=horas,estado=request.form.get('estado','Planificada'),observacion=request.form.get('observacion','').strip())
        db.session.add(a); db.session.commit(); audit('Actividad planificada',f'{ot.numero}: {tarea}'); return redirect(url_for('planificacion',mes=fi.month,anio=fi.year))
    import calendar
    ndias=calendar.monthrange(anio,mes)[1]; dias=list(range(1,ndias+1)); inicio=date(anio,mes,1); fin=date(anio,mes,ndias)
    actividades=PlanActividad.query.filter(PlanActividad.fecha_inicio<=fin,PlanActividad.fecha_fin>=inicio).order_by(PlanActividad.fecha_inicio,PlanActividad.id).all()
    usuarios=operarios_planificables()
    usuarios_pendientes=Usuario.query.filter_by(activo=True,disponible_planificacion=False).order_by(Usuario.nombre).all()
    # La aprobación puede constar en la OT o en su cotización activa.
    # No limitar el listado a una sola coincidencia literal ni a la primera OT.
    aprobadas=[]
    for ot in OT.query.order_by(OT.fecha.desc()).all():
        q=Cotizacion.query.filter(Cotizacion.ot_id==ot.id, Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).first()
        estado_cot=(q.estado or '').strip().casefold() if q else ''
        estado_ot=(ot.aprobacion or '').strip().casefold()
        # Una cotización expresamente rechazada no puede programarse.
        if estado_cot=='rechazada':
            continue
        if estado_cot in ('aprobada','aprobado') or (q and estado_ot in ('aprobada','aprobado')):
            aprobadas.append(ot)
    carga={u.id:{d:0 for d in dias} for u in usuarios}
    for a in actividades:
        d=max(a.fecha_inicio,inicio)
        while d<=min(a.fecha_fin,fin):
            if d.day in carga.get(a.usuario_id,{}): carga[a.usuario_id][d.day]+=a.horas_dia or 0
            d+=timedelta(days=1)
    # Resumen compacto: agrupar por OT, componente, etapa y responsable.
    # Conservar asignaciones independientes para edición y auditoría.
    resumen={}
    for a in sorted(actividades, key=lambda x:(x.ot.numero if x.ot else '', x.componente.nombre if x.componente else '', x.fecha_inicio, x.id)):
        ot_key=(a.ot_id, a.ot.numero if a.ot else 'OT sin número')
        comp_key=(a.ot_componente_id or 0, a.componente.nombre if a.componente else 'Sin componente')
        nombre_maestro=(a.reparacion.nombre or '').strip() if a.reparacion else ''
        nombre_manual=(a.tarea or '').strip()
        # No atribuir automáticamente tareas antiguas a una reparación distinta.
        nombre=nombre_maestro or nombre_manual or f'Actividad histórica #{a.id} (sin descripción)'
        etapa_key=(a.reparacion_maestra_id or 0, nombre)
        por_etapa=resumen.setdefault(ot_key,{}).setdefault(comp_key,{}).setdefault(etapa_key,{})
        grupo=por_etapa.setdefault(a.usuario_id,{'responsable':a.responsable,'asignaciones':[],'horas_por_dia':{},'total_hh':0})
        grupo['asignaciones'].append(a)
        dia=max(a.fecha_inicio,inicio)
        while dia<=min(a.fecha_fin,fin):
            grupo['horas_por_dia'][dia.day]=grupo['horas_por_dia'].get(dia.day,0)+(a.horas_dia or 0)
            grupo['total_hh']+=(a.horas_dia or 0)
            dia+=timedelta(days=1)
    return render_template('planificacion.html',mes=mes,anio=anio,dias=dias,actividades=actividades,usuarios=usuarios,usuarios_pendientes=usuarios_pendientes,ots=aprobadas,carga=carga,inicio=inicio,resumen=resumen)


@app.route('/planificacion/resumen')
@permiso_required('planificacion')
def planificacion_resumen():
    import calendar
    hoy=date.today()
    try:
        mes=int(request.args.get('mes',hoy.month)); anio=int(request.args.get('anio',hoy.year))
        inicio=date(anio,mes,1); fin=date(anio,mes,calendar.monthrange(anio,mes)[1])
    except (ValueError,OverflowError):
        abort(400)
    actividades=PlanActividad.query.filter(PlanActividad.fecha_inicio<=fin,PlanActividad.fecha_fin>=inicio).all()
    por_persona={}; ot_ids=set(); total=0; por_estado={}
    for actividad in actividades:
        d=max(actividad.fecha_inicio,inicio); hasta=min(actividad.fecha_fin,fin)
        if d>hasta: continue
        hh=((hasta-d).days+1)*float(actividad.horas_dia or 0)
        total+=hh;ot_ids.add(actividad.ot_id)
        estado=actividad.estado or 'Sin estado';por_estado[estado]=por_estado.get(estado,0)+1
        u=actividad.responsable
        clave=actividad.usuario_id
        item=por_persona.setdefault(clave,{'nombre':u.nombre if u else 'Sin responsable','sigla':u.sigla if u else '—','color':u.color if u and u.color else '#64748b','hh':0,'dias':{}})
        item['hh']+=hh
        while d<=hasta:
            item['dias'][d.day]=item['dias'].get(d.day,0)+float(actividad.horas_dia or 0)
            d+=timedelta(days=1)
    personas=sorted(por_persona.values(),key=lambda x:x['nombre'])
    return render_template('planificacion_resumen.html',mes=mes,anio=anio,total=total,ots=len(ot_ids),actividades=len(actividades),personas=personas,por_estado=por_estado)


@app.route('/planificacion/gantt')
@permiso_required('planificacion')
def planificacion_gantt():
    import calendar
    hoy=date.today()
    try:
        mes=int(request.args.get('mes') or hoy.month)
        anio=int(request.args.get('anio') or hoy.year)
        inicio=date(anio,mes,1)
    except (ValueError,TypeError): abort(400)
    fin=date(anio,mes,calendar.monthrange(anio,mes)[1])
    trabajadores=operarios_planificables()
    aprobadas=[]
    for ot in OT.query.order_by(OT.fecha.desc(),OT.id.desc()).all():
        q=Cotizacion.query.filter(Cotizacion.ot_id==ot.id,Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).first()
        ec=(q.estado or '').strip().casefold() if q else ''
        eo=(ot.aprobacion or '').strip().casefold()
        if ec=='rechazada': continue
        if ec in ('aprobada','aprobado') or (q and eo in ('aprobada','aprobado')):
            aprobadas.append(ot)
    lista=[]
    for ot in aprobadas:
        componentes=[]
        for c in ot.componentes_detalle:
            if (c.condicion_ingreso or '').strip().casefold()!='recibido': continue
            etapas=[]
            maestros=ReparacionMaestra.query.filter_by(componente_modelo_id=c.componente_modelo_id,activo=True).order_by(ReparacionMaestra.secuencia,ReparacionMaestra.id).all() if c.componente_modelo_id else []
            for r in maestros:
                na=PlanNoAplica.query.filter_by(ot_id=ot.id,ot_componente_id=c.id,reparacion_maestra_id=r.id).first() is not None
                acts=asignaciones_reparacion(ot.id,c.id,r)
                etapas.append({'nombre':r.nombre,'rep_id':r.id,'actividades':acts,'estimadas':float(r.hh_estimadas or 0),'no_aplica':na})
            # Mantener visibles actividades históricas o manuales no presentes en el maestro.
            ids={a.id for e in etapas for a in e['actividades']}
            for a in PlanActividad.query.filter_by(ot_id=ot.id,ot_componente_id=c.id).all():
                if a.id not in ids: etapas.append({'nombre':a.tarea or 'Actividad histórica','rep_id':None,'actividades':[a],'estimadas':0})
            if etapas: componentes.append({'nombre':c.nombre,'id':c.id,'etapas':etapas})
        if componentes: lista.append({'id':ot.id,'numero':ot.numero,'componentes':componentes})
    import json as _json
    def ajson(a):
        return {'id':a.id,'worker':a.usuario_id,'start':a.fecha_inicio.isoformat(),'end':a.fecha_fin.isoformat(),'hours':a.horas_dia,'sigla':a.responsable.sigla if a.responsable else '?','liberada':bool(PlanLiberacion.query.filter_by(asignacion_id=a.id,consumido_fecha=None).first())}
    data=[{'id':o['id'],'numero':o['numero'],'components':[{'id':c['id'],'nombre':c['nombre'],'tasks':[{'name':e['nombre'],'rep':e['rep_id'],'assignments':[ajson(a) for a in e['actividades']],'no_aplica':e.get('no_aplica',False)} for e in c['etapas']]} for c in o['componentes']]} for o in lista]
    workers=[{'id':u.id,'name':u.nombre,'sigla':u.sigla or u.nombre[:2],'color':(u.color or '#64748b')} for u in trabajadores]
    return render_template('planificacion_gantt.html',gantt_data=data,gantt_workers=workers,mes=mes,anio=anio)


@app.route('/planificacion/gantt/pdf/<tipo>')
@permiso_required('planificacion')
def planificacion_gantt_pdf(tipo):
    """Exporta programación vigente desde PostgreSQL, no desde datos de demostración."""
    from io import BytesIO
    import calendar
    from flask import send_file
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    if tipo not in ('mensual','trabajador'): abort(404)
    try:
        mes=int(request.args.get('mes',date.today().month)); anio=int(request.args.get('anio',date.today().year))
        inicio=date(anio,mes,1); fin=date(anio,mes,calendar.monthrange(anio,mes)[1])
        trabajador_id=int(request.args.get('trabajador_id') or 0)
    except (TypeError,ValueError): abort(400)
    usuario=db.session.get(Usuario,trabajador_id) if tipo=='trabajador' else None
    if tipo=='trabajador' and not usuario: abort(400)
    buffer=BytesIO(); page=landscape(A3 if tipo=='mensual' else A4)
    doc=SimpleDocTemplate(buffer,pagesize=page,leftMargin=20,rightMargin=20,topMargin=25,bottomMargin=25)
    styles=getSampleStyleSheet(); story=[]
    titulo='PLANIFICACIÓN MENSUAL GENERAL' if tipo=='mensual' else 'PLANIFICACIÓN INDIVIDUAL · '+usuario.nombre
    story.append(Paragraph('INGEPRO · DIVISIÓN MAESTRANZA',styles['Heading2']))
    story.append(Paragraph(f'{titulo} · {mes:02d}/{anio}',styles['Heading3']))
    story.append(Spacer(1,9))
    if tipo=='mensual':
        encabezado=['OT / COMPONENTE / ETAPA','ESTADO']+[str(d) for d in range(1,fin.day+1)]
        filas=[encabezado]
        ots=OT.query.order_by(OT.numero).all()
        for ot in ots:
            q=Cotizacion.query.filter(Cotizacion.ot_id==ot.id,Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).first()
            if not q or ((q.estado or '').strip().casefold() not in ('aprobada','aprobado') and (ot.aprobacion or '').strip().casefold() not in ('aprobada','aprobado')):continue
            for c in ot.componentes_detalle:
                if (c.condicion_ingreso or '').strip().casefold()!='recibido' or not c.componente_modelo_id:continue
                reps=ReparacionMaestra.query.filter_by(componente_modelo_id=c.componente_modelo_id,activo=True).order_by(ReparacionMaestra.secuencia,ReparacionMaestra.id).all()
                if not reps:continue
                filas.append([f'{ot.numero} / {c.nombre}'[:49],'']+['']*fin.day)
                for rep in reps:
                    na=PlanNoAplica.query.filter_by(ot_id=ot.id,ot_componente_id=c.id,reparacion_maestra_id=rep.id).first() is not None
                    row=[rep.nombre[:49],'N/A' if na else 'Aplica']+['']*fin.day
                    if not na:
                        for a in asignaciones_reparacion(ot.id,c.id,rep):
                            for d in range(1,fin.day+1):
                                fecha=date(anio,mes,d)
                                if a.fecha_inicio<=fecha<=a.fecha_fin:
                                    v=f'{(a.responsable.sigla or "?") if a.responsable else "?"} {float(a.horas_dia or 0):g}h'
                                    row[d+1]=(row[d+1]+' / '+v).strip(' /')
                    filas.append(row)
        if len(filas)==1:filas.append(['Sin reparaciones aprobadas para mostrar','']+['']*fin.day)
        ancho=(page[0]-40-260-43)/fin.day
        tabla=Table(filas,colWidths=[260,43]+[ancho]*fin.day,repeatRows=1,hAlign='LEFT')
        tabla.setStyle(TableStyle([('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),5),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e0e3e8')),('GRID',(0,0),(-1,-1),.2,colors.HexColor('#d6dbe1')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8fafc')])]))
    else:
        acts=PlanActividad.query.filter(PlanActividad.usuario_id==usuario.id,PlanActividad.fecha_inicio<=fin,PlanActividad.fecha_fin>=inicio).order_by(PlanActividad.fecha_inicio,PlanActividad.ot_id).all()
        filas=[['FECHA','OT','COMPONENTE','ETAPA','HORAS']]
        total=0.0
        for a in acts:
            for d in range(1,fin.day+1):
                fecha=date(anio,mes,d)
                if a.fecha_inicio<=fecha<=a.fecha_fin:
                    filas.append([fecha.strftime('%d/%m/%Y'),a.ot.numero if a.ot else '',a.componente.nombre[:28] if a.componente else '',a.tarea[:43],f'{float(a.horas_dia or 0):g}'])
                    total+=float(a.horas_dia or 0)
        filas.append(['TOTAL','','','',f'{total:g}'])
        tabla=Table(filas,colWidths=[85,95,165,255,55],repeatRows=1,hAlign='LEFT')
        tabla.setStyle(TableStyle([('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e0e3e8')),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#d6dbe1')),('FONTSIZE',(0,0),(-1,-1),8),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8fafc')])]))
    story.append(tabla);doc.build(story);buffer.seek(0)
    return send_file(buffer,mimetype='application/pdf',as_attachment=True,download_name=f'INGEPRO_{tipo}_{anio}_{mes:02d}.pdf')


@app.route('/api/planificacion/asignacion/<int:aid>/liberar',methods=['POST'])
@permiso_required('planificacion')
def plan_liberar(aid):
    if not usuario_actual().es_admin: return jsonify(error='Solo el administrador puede liberar asignaciones.'),403
    asignacion=db.session.get(PlanActividad,aid)
    if not asignacion: return jsonify(error='Asignación no encontrada.'),404
    permiso=PlanLiberacion.query.filter_by(asignacion_id=aid).first()
    if not permiso:
        permiso=PlanLiberacion(asignacion_id=aid,liberado_por_id=usuario_actual().id)
        db.session.add(permiso)
    permiso.liberado_por_id=usuario_actual().id
    permiso.liberado_fecha=datetime.utcnow()
    permiso.consumido_por_id=None
    permiso.consumido_fecha=None
    db.session.commit()
    audit('Asignación liberada por administrador',f'Asignación {aid}, OT {asignacion.ot_id}, administrador {usuario_actual().username}')
    return jsonify(ok=True,mensaje='Asignación liberada para una modificación del supervisor.')

@app.route('/api/planificacion/gantt/rapido', methods=['POST'])
@permiso_required('planificacion')
def api_gantt_rapido():
    """Edición por día sin salir de la carta Gantt; conserva los demás días de rangos."""
    data=request.get_json(silent=True) or {}
    try:
        dia=date.fromisoformat(str(data['fecha']))
        ot_id=int(data['ot_id']); comp_id=int(data['componente_id'])
        rep_id=int(data['reparacion_id']) if data.get('reparacion_id') else None
        aid=int(data['asignacion_id']) if data.get('asignacion_id') else None
        accion=data.get('accion','guardar')
        uid=int(data['usuario_id']) if accion=='guardar' else None
        horas=float(data['horas']) if accion=='guardar' else None
    except (KeyError,ValueError,TypeError): return jsonify(error='Datos inválidos.'),400
    if accion not in ('guardar','eliminar'): return jsonify(error='Acción no válida.'),400
    ot=db.session.get(OT,ot_id); comp=db.session.get(OTComponente,comp_id)
    if not ot or not comp or comp.ot_id!=ot_id or (comp.condicion_ingreso or '').casefold()!='recibido':
        return jsonify(error='OT o componente no disponible.'),400
    # V7.56: supervisor puede CREAR, pero no alterar ni eliminar una asignación.
    # La validación es del servidor, no depende de controles HTML.
    if es_supervisor():
        if accion!='guardar':
            return jsonify(error='El supervisor no puede eliminar asignaciones.'),403
        if aid:
            permiso=PlanLiberacion.query.filter_by(asignacion_id=aid,consumido_fecha=None).first()
            if not permiso:
                return jsonify(error='Asignación bloqueada. Solicite su liberación al administrador.'),403
    original=db.session.get(PlanActividad,aid) if aid else None
    if aid and (not original or original.ot_id!=ot_id or original.ot_componente_id!=comp_id or not(original.fecha_inicio<=dia<=original.fecha_fin)):
        return jsonify(error='La asignación ya no corresponde a esta casilla.'),409
    rep=db.session.get(ReparacionMaestra,rep_id) if rep_id else None
    if not original and (not rep or not rep.activo or rep.componente_modelo_id!=comp.componente_modelo_id):
        return jsonify(error='La etapa no admite una asignación nueva.'),400
    if not original and PlanNoAplica.query.filter_by(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id).first():
        return jsonify(error='Esta reparación está marcada No aplica.'),400
    if accion=='eliminar' and not original: return jsonify(error='No hay asignación que eliminar.'),400
    usuario=db.session.get(Usuario,uid) if uid else None
    if accion=='guardar':
        error=validar_cupo(usuario,dia,dia,horas,excluir_id=aid)
        if error: return jsonify(error=error),400
        # No duplicar el mismo trabajador en una etapa el mismo día.
        similares=PlanActividad.query.filter(PlanActividad.ot_id==ot_id,PlanActividad.ot_componente_id==comp_id,PlanActividad.usuario_id==uid,PlanActividad.fecha_inicio<=dia,PlanActividad.fecha_fin>=dia).all()
        for a in similares:
            if a.id==aid: continue
            if (rep_id and a.reparacion_maestra_id==rep_id) or (original and a.tarea==original.tarea):
                return jsonify(error='Este trabajador ya está asignado a esta etapa ese día.'),409
        # La estimación del maestro no limita las horas planificadas.
        # validar_cupo() controla el máximo diario y evita sobreasignaciones.
    # Una asignación de varios días se divide para cambiar únicamente la fecha elegida.
    if original:
        start,end=original.fecha_inicio,original.fecha_fin
        if start!=end:
            def clonar(fi,ff):
                return PlanActividad(ot_id=original.ot_id,ot_componente_id=original.ot_componente_id,reparacion_maestra_id=original.reparacion_maestra_id,usuario_id=original.usuario_id,tarea=original.tarea,fecha_inicio=fi,fecha_fin=ff,horas_dia=original.horas_dia,estado=original.estado,observacion=original.observacion)
            if start<dia: db.session.add(clonar(start,dia-timedelta(days=1)))
            if dia<end: db.session.add(clonar(dia+timedelta(days=1),end))
            original.fecha_inicio=dia;original.fecha_fin=dia
    if accion=='eliminar': db.session.delete(original)
    elif original:
        original.usuario_id=uid;original.horas_dia=horas
    else:
        db.session.add(PlanActividad(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id,usuario_id=uid,tarea=rep.nombre,fecha_inicio=dia,fecha_fin=dia,horas_dia=horas,estado='Planificada'))
        if not OTReparacionSeleccionada.query.filter_by(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id).first():
            db.session.add(OTReparacionSeleccionada(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id))
    if es_supervisor() and original:
        permiso.consumido_por_id=usuario_actual().id
        permiso.consumido_fecha=datetime.utcnow()
    db.session.commit()
    if es_supervisor() and original:
        audit('Asignación modificada tras liberación',f'Asignación {aid}; administrador liberador {permiso.liberado_por_id}; supervisor {usuario_actual().username}')
    audit('Gantt edición rápida',f'OT {ot.numero} / {dia.isoformat()} / {accion}')
    return jsonify(ok=True)


@app.route('/api/planificacion/ot/<int:ot_id>')
@permiso_required('planificacion')
def api_plan_ot(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    data=[]
    for c in ot.componentes_detalle:
        if (c.condicion_ingreso or '').strip().casefold()!='recibido':
            continue
        selected={x.reparacion_maestra_id for x in OTReparacionSeleccionada.query.filter_by(ot_id=ot.id,ot_componente_id=c.id).all()}
        reps=[]
        if c.componente_modelo_id:
            for r in ReparacionMaestra.query.filter_by(componente_modelo_id=c.componente_modelo_id,activo=True).order_by(ReparacionMaestra.secuencia,ReparacionMaestra.id).all():
                asignaciones=asignaciones_reparacion(ot.id,c.id,r)
                programadas=sum(((a.fecha_fin-a.fecha_inicio).days+1)*(a.horas_dia or 0) for a in asignaciones)
                no_aplica=PlanNoAplica.query.filter_by(ot_id=ot.id,ot_componente_id=c.id,reparacion_maestra_id=r.id).first() is not None
                estimadas=float(r.hh_estimadas or 0)
                restante=max(0,estimadas-programadas) if estimadas>0 else None
                estado='No aplica' if no_aplica else ('Programada' if asignaciones else 'Pendiente')
                reps.append({'id':r.id,'nombre':r.nombre,'hh':estimadas,'seleccionada':r.id in selected,'programadas':programadas,'restante':restante,'estado':estado})
        data.append({'id':c.id,'nombre':c.nombre,'vinculado':bool(c.componente_modelo_id),'reparaciones':reps})
    return jsonify(data)

@app.route('/api/planificacion/reparacion/estado',methods=['POST'])
@permiso_required('planificacion')
def plan_reparacion_estado():
    payload=request.get_json(silent=True) if request.is_json else request.form
    payload=payload or {}
    try:
        ot_id=int(payload.get('ot_id') or 0)
        comp_id=int(payload.get('ot_componente_id') or 0)
        rep_id=int(payload.get('reparacion_maestra_id') or 0)
    except (ValueError,TypeError): abort(400)
    ot=db.session.get(OT,ot_id); comp=db.session.get(OTComponente,comp_id); rep=db.session.get(ReparacionMaestra,rep_id)
    if not ot or not comp or comp.ot_id!=ot_id or not rep or rep.componente_modelo_id!=comp.componente_modelo_id: abort(400)
    accion=payload.get('accion')
    registro=PlanNoAplica.query.filter_by(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id).first()
    if accion=='no_aplica':
        if asignaciones_reparacion(ot_id,comp_id,rep):
            return jsonify({'error':'La reparación tiene horas programadas. Elimine primero sus asignaciones.'}),409
        if not registro:
            db.session.add(PlanNoAplica(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id,usuario=str(session.get('username',''))))
    elif accion=='revertir':
        if registro: db.session.delete(registro)
    else: abort(400)
    db.session.commit()
    audit('Estado de reparación en planificación',f'{ot.numero} · {rep.nombre}: {accion}')
    return jsonify({'ok':True})

@app.route('/planificacion/habilitar-operario', methods=['POST'])
@permiso_required('usuarios')
def plan_habilitar_operario():
    try: uid=int(request.form.get('usuario_id') or 0)
    except ValueError: abort(400)
    u=db.session.get(Usuario,uid) or abort(404)
    if not u.activo: abort(400)
    # Solo usuarios internos; ContactoCliente es una entidad distinta.
    u.disponible_planificacion=True
    db.session.commit()
    audit('Operario habilitado para planificación',u.username)
    flash(f'{u.nombre} habilitado como responsable de planificación.')
    return redirect(url_for('planificacion'))

@app.route('/api/planificacion/ot/<int:ot_id>/estado')
@permiso_required('planificacion')
def api_plan_ot_estado(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    detalles=[]
    for c in ot.componentes_detalle:
        if (c.condicion_ingreso or '').strip().lower()!='recibido': continue
        total=ReparacionMaestra.query.filter_by(componente_modelo_id=c.componente_modelo_id,activo=True).count() if c.componente_modelo_id else 0
        elegidas=OTReparacionSeleccionada.query.filter_by(ot_componente_id=c.id).count()
        detalles.append({'id':c.id,'nombre':c.nombre,'vinculado':bool(c.componente_modelo_id),'maestro':total,'seleccionadas':elegidas})
    return jsonify({'componentes':detalles,'ot_url':url_for('ver_ot',ot_id=ot.id),'diagnostico_url':url_for('ot_reparaciones',ot_id=ot.id)})

@app.route('/planificacion/<int:aid>/editar',methods=['GET','POST'])
@permiso_required('planificacion')
def plan_editar(aid):
    if es_supervisor(): abort(403)
    a=db.session.get(PlanActividad,aid) or abort(404)
    if request.method=='POST':
        usuario=db.session.get(Usuario,int(request.form.get('usuario_id') or 0))
        if not usuario or not usuario.activo or not usuario.disponible_planificacion: abort(400)
        try:
            fi=date.fromisoformat(request.form['fecha_inicio']); ff=date.fromisoformat(request.form['fecha_fin'])
            horas=float(request.form['horas_dia'])
        except (ValueError,KeyError): abort(400)
        error=validar_cupo(usuario,fi,ff,horas,excluir_id=aid)
        if error: flash(error); return redirect(url_for('plan_editar',aid=aid))
        tarea=request.form.get('tarea','').strip()
        if not tarea: flash('La tarea es obligatoria.'); return redirect(url_for('plan_editar',aid=aid))
        a.usuario_id=usuario.id; a.fecha_inicio=fi; a.fecha_fin=ff; a.horas_dia=horas
        a.tarea=tarea; a.estado=request.form.get('estado','Planificada'); a.observacion=request.form.get('observacion','').strip()
        db.session.commit(); audit('Planificación modificada',f'{a.ot.numero}: {a.tarea}')
        flash('Actividad actualizada correctamente.')
        return redirect(url_for('planificacion',mes=fi.month,anio=fi.year))
    usuarios=operarios_planificables()
    return render_template('planificacion_editar.html',a=a,usuarios=usuarios)

@app.route('/planificacion/<int:aid>/duplicar',methods=['POST'])
@permiso_required('planificacion')
def plan_duplicar(aid):
    if es_supervisor(): abort(403)
    a=db.session.get(PlanActividad,aid) or abort(404)
    u=db.session.get(Usuario,int(request.form.get('usuario_id') or 0))
    if not u or not u.activo or not u.disponible_planificacion: abort(400)
    error=validar_cupo(u,a.fecha_inicio,a.fecha_fin,a.horas_dia)
    if error: flash(error); return redirect(url_for('planificacion',mes=a.fecha_inicio.month,anio=a.fecha_inicio.year))
    copia=PlanActividad(ot_id=a.ot_id,ot_componente_id=a.ot_componente_id,reparacion_maestra_id=a.reparacion_maestra_id,
        usuario_id=u.id,tarea=a.tarea,fecha_inicio=a.fecha_inicio,fecha_fin=a.fecha_fin,
        horas_dia=a.horas_dia,estado=a.estado,observacion=a.observacion)
    db.session.add(copia); db.session.commit(); audit('Segundo responsable planificado',f'{a.ot.numero}: {a.tarea} / {u.nombre}')
    flash('Responsable adicional asignado. Puede editar sus fechas y horas por separado.')
    return redirect(url_for('planificacion',mes=a.fecha_inicio.month,anio=a.fecha_inicio.year))

@app.route('/planificacion/<int:aid>/eliminar',methods=['POST'])
@permiso_required('planificacion')
def plan_eliminar(aid):
    if es_supervisor(): abort(403)
    a=db.session.get(PlanActividad,aid) or abort(404); m=a.fecha_inicio.month; y=a.fecha_inicio.year; db.session.delete(a); db.session.commit(); audit('Asignación eliminada',f'{a.ot.numero}: {a.tarea}'); flash('Asignación eliminada; la OT y el diagnóstico permanecen intactos.'); return redirect(url_for('planificacion',mes=m,anio=y))

def cotizacion_activa_ot(ot_id):
    return Cotizacion.query.filter(Cotizacion.ot_id==ot_id, Cotizacion.estado!='Anulada').order_by(Cotizacion.id.desc()).first()

def cargar_cotizacion_form(q, form):
    q.contacto=form.get('contacto','').strip(); q.plazo_pago=form.get('plazo_pago','30 días').strip()
    q.referencia=form.get('referencia','').strip(); q.notas=form.get('notas','').strip()
    q.plazo_ejecucion_cantidad=int(form.get('plazo_ejecucion_cantidad') or 0)
    q.plazo_ejecucion_unidad=form.get('plazo_ejecucion_unidad','Días hábiles')
    q.inicio_plazo=form.get('inicio_plazo','Recepción de OC / aprobación').strip()
    try: q.descuento_pct=max(0,min(float(form.get('descuento_pct') or 0),100))
    except: q.descuento_pct=0
    vh=form.get('validez_hasta'); q.validez_hasta=datetime.strptime(vh,'%Y-%m-%d').date() if vh else None
    q.lineas.clear(); db.session.flush()
    descs=form.getlist('descripcion[]'); cants=form.getlist('cantidad[]'); precios=form.getlist('precio[]')
    for i,desc in enumerate(descs):
        if not desc.strip(): continue
        try: cant=float(cants[i] or 1); precio=int(float(precios[i] or 0))
        except: cant,precio=1,0
        q.lineas.append(LineaCotizacion(descripcion=desc.strip(),cantidad=cant,precio_unitario=precio,orden=i))

@app.route('/cotizaciones')
@permiso_required('cotizaciones')
def cotizaciones():
    qs=Cotizacion.query.order_by(Cotizacion.id.desc()).all()
    return render_template('cotizaciones.html',cotizaciones=qs)

@app.route('/ot/<int:ot_id>/cotizacion/nueva', methods=['GET','POST'])
@permiso_required('cotizaciones')
def cotizacion_nueva(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    activa=cotizacion_activa_ot(ot.id)
    if activa:
        flash(f'Esta OT ya tiene una cotización activa: {activa.numero}.')
        return redirect(url_for('ver_cotizacion',qid=activa.id))
    if not ((ot.diagnostico or '').strip() or ot.hallazgos):
        flash('Debes registrar un diagnóstico antes de crear la cotización.')
        return redirect(url_for('diagnostico_ot',ot_id=ot.id))
    if request.method=='POST':
        q=Cotizacion(numero=next_quote_number(),ot_id=ot.id,cliente=ot.cliente,creada_por=usuario_actual().nombre)
        cargar_cotizacion_form(q,request.form); db.session.add(q); db.session.commit()
        ot.cotizacion=q.numero
        if request.form.get('accion')=='enviar': q.estado='Presupuesto enviado'; ot.estado='Cotización enviada'
        db.session.commit(); audit('Cotización creada',f'{q.numero} / {ot.numero}'); flash(f'{q.numero} creada correctamente')
        return redirect(url_for('ver_cotizacion',qid=q.id))
    sugerencias=[f'REPARACIÓN {x.strip().upper()}' for x in ot.componente.split(',')] if ot.componente and ot.componente!='Bomba completa' else ['REPARACIÓN BOMBA COMPLETA']
    return render_template('cotizacion_form.html',ot=ot,sugerencias=sugerencias,q=None,modo='nueva')

@app.route('/cotizacion/<int:qid>')
@permiso_required('cotizaciones')
def ver_cotizacion(qid):
    q=db.session.get(Cotizacion,qid) or abort(404)
    return render_template('cotizacion_ver.html',q=q,ot=db.session.get(OT,q.ot_id))

@app.route('/cotizacion/<int:qid>/editar',methods=['GET','POST'])
@permiso_required('cotizaciones')
def editar_cotizacion(qid):
    q=db.session.get(Cotizacion,qid) or abort(404); ot=db.session.get(OT,q.ot_id)
    if q.estado=='Anulada': flash('Una cotización anulada no puede modificarse.'); return redirect(url_for('ver_cotizacion',qid=q.id))
    if request.method=='POST':
        cargar_cotizacion_form(q,request.form)
        if request.form.get('accion')=='enviar': q.estado='Presupuesto enviado'; ot.estado='Cotización enviada'
        db.session.commit(); audit('Cotización editada',q.numero); flash('Cotización actualizada correctamente.')
        return redirect(url_for('ver_cotizacion',qid=q.id))
    return render_template('cotizacion_form.html',ot=ot,sugerencias=[],q=q,modo='editar')

@app.route('/cotizacion/<int:qid>/anular',methods=['POST'])
@permiso_required('cotizaciones')
def anular_cotizacion(qid):
    q=db.session.get(Cotizacion,qid) or abort(404); motivo=request.form.get('motivo','').strip()
    if not motivo: flash('Debes indicar el motivo de anulación.'); return redirect(url_for('ver_cotizacion',qid=q.id))
    q.estado='Anulada'; q.anulada_motivo=motivo; q.anulada_fecha=datetime.utcnow(); q.anulada_por=usuario_actual().nombre
    ot=db.session.get(OT,q.ot_id); ot.cotizacion=None; ot.aprobacion='Pendiente'
    db.session.commit(); audit('Cotización anulada',f'{q.numero}: {motivo}'); flash(f'{q.numero} quedó ANULADA. Ya puedes crear una nueva cotización para la OT.')
    return redirect(url_for('ver_ot',ot_id=ot.id))

@app.route('/cotizacion/<int:qid>/estado', methods=['POST'])
@permiso_required('cotizaciones')
def cotizacion_estado(qid):
    q=db.session.get(Cotizacion,qid) or abort(404); nuevo=request.form.get('estado')
    if q.estado=='Anulada': return redirect(url_for('ver_cotizacion',qid=q.id))
    if nuevo in ['Presupuesto','Presupuesto enviado','Aprobada','Rechazada']:
        q.estado=nuevo; ot=db.session.get(OT,q.ot_id)
        if nuevo=='Aprobada':
            if not ((ot.diagnostico or '').strip() or ot.hallazgos):
                flash('No se puede aprobar una reparación sin diagnóstico registrado.')
                return redirect(url_for('ver_cotizacion',qid=q.id))
            ot.aprobacion='Aprobada'; ot.estado='Aprobada / Reparación'; q.fecha_aprobacion=datetime.utcnow(); q.fecha_comprometida=sumar_plazo(date.today(),q.plazo_ejecucion_cantidad,q.plazo_ejecucion_unidad); ot.fecha_comprometida=q.fecha_comprometida
        elif nuevo=='Rechazada': ot.aprobacion='Rechazada'; ot.estado='Cotización rechazada'
        elif nuevo=='Presupuesto enviado': ot.estado='Cotización enviada'
        db.session.commit(); audit('Estado cotización',f'{q.numero}: {nuevo}')
    return redirect(url_for('ver_cotizacion',qid=q.id))

@app.route('/cotizacion/<int:qid>/pdf')
@permiso_required('cotizaciones')
def cotizacion_pdf(qid):
    q=db.session.get(Cotizacion,qid) or abort(404); ot=db.session.get(OT,q.ot_id)
    buf=BytesIO(); doc=SimpleDocTemplate(buf,pagesize=A4,rightMargin=28,leftMargin=28,topMargin=28,bottomMargin=28)
    styles=getSampleStyleSheet(); story=[]
    title=ParagraphStyle('t',parent=styles['Title'],fontSize=22,textColor=colors.HexColor('#0b6596'),alignment=TA_RIGHT)
    logo_path=os.path.join(app.root_path,'static','img','logo-ingepro.png')
    logo=Image(logo_path,width=170,height=70) if os.path.exists(logo_path) else Paragraph('<b>INGEPRO</b>',styles['Heading1'])
    cab=Table([[logo,Paragraph(f'<b>DIVISIÓN MAESTRANZA</b><br/><font size="16">Presupuesto {q.numero}</font>',title)]],colWidths=[220,290]); cab.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LINEBELOW',(0,0),(-1,-1),2,colors.HexColor('#0b6596')),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story += [cab,Spacer(1,14)]
    meta=[[Paragraph('<b>Cliente</b>',styles['BodyText']),q.cliente],[Paragraph('<b>OT</b>',styles['BodyText']),ot.numero],[Paragraph('<b>Equipo</b>',styles['BodyText']),ot.equipo],[Paragraph('<b>Fecha</b>',styles['BodyText']),q.fecha.strftime('%d/%m/%Y')],[Paragraph('<b>Validez</b>',styles['BodyText']),q.validez_hasta.strftime('%d/%m/%Y') if q.validez_hasta else '-'],[Paragraph('<b>Tiempo de ejecución</b>',styles['BodyText']),f'{q.plazo_ejecucion_cantidad or "-"} {q.plazo_ejecucion_unidad or ""}']]
    mt=Table(meta,colWidths=[120,390]); mt.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.lightgrey),('BACKGROUND',(0,0),(0,-1),colors.HexColor('#eef3f7')),('PADDING',(0,0),(-1,-1),6)])); story += [mt,Spacer(1,14)]
    data=[['DESCRIPCIÓN','CANT.','PRECIO UNITARIO','SUBTOTAL']]
    for l in q.lineas: data.append([Paragraph(l.descripcion,styles['BodyText']),f'{l.cantidad:g}',f'$ {l.precio_unitario:,.0f}'.replace(',','.'),f'$ {l.cantidad*l.precio_unitario:,.0f}'.replace(',','.')])
    t=Table(data,colWidths=[270,55,95,90],repeatRows=1); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0b6596')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.4,colors.grey),('ALIGN',(1,1),(-1,-1),'RIGHT'),('PADDING',(0,0),(-1,-1),6)])); story += [t,Spacer(1,10)]
    totals=[['Subtotal',f'$ {q.subtotal:,.0f}'.replace(',','.')], [f'Descuento {q.descuento_pct or 0:g}%',f'- $ {q.descuento:,.0f}'.replace(',','.')],['Neto',f'$ {q.neto:,.0f}'.replace(',','.')],['IVA 19%',f'$ {q.iva:,.0f}'.replace(',','.')],['TOTAL',f'$ {q.total:,.0f}'.replace(',','.')]]
    tt=Table(totals,colWidths=[120,120],hAlign='RIGHT'); tt.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.grey),('ALIGN',(1,0),(1,-1),'RIGHT'),('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#0b6596')),('TEXTCOLOR',(0,-1),(-1,-1),colors.white),('PADDING',(0,0),(-1,-1),6)])); story += [tt,Spacer(1,14)]
    story += [Paragraph(f'<b>Condiciones de pago:</b> {q.plazo_pago}',styles['BodyText']),Paragraph(f'<b>Inicio del plazo:</b> {q.inicio_plazo}',styles['BodyText'])]
    if q.referencia: story.append(Paragraph(f'<b>Referencia:</b> {q.referencia}',styles['BodyText']))
    if q.notas: story.append(Paragraph(f'<b>Notas:</b> {q.notas}',styles['BodyText']))
    if q.estado=='Anulada': story.append(Paragraph('<b>COTIZACIÓN ANULADA</b>',styles['Heading2']))
    doc.build(story); buf.seek(0)
    audit('PDF cotización generado',q.numero)
    return send_file(buf,mimetype='application/pdf',as_attachment=True,download_name=f'{q.numero}.pdf')

# Limpieza administrativa de datos transaccionales. Nunca modifica tablas maestras.
TABLAS_LIMPIEZA=(LineaCotizacion,Cotizacion,PlanNoAplica,PlanActividad,OTReparacionSeleccionada,DiagnosticoHallazgo,OTComponente,OT)

def conteos_limpieza():
    return [(cls.__tablename__, db.session.query(cls).count()) for cls in TABLAS_LIMPIEZA]

def respaldo_completo():
    url=app.config['SQLALCHEMY_DATABASE_URI']
    if url.startswith('sqlite:'):
        import sqlalchemy
        ruta=db.engine.url.database
        with tempfile.NamedTemporaryFile(suffix='.sqlite3',delete=False) as f: destino=f.name
        try:
            src=sqlite3.connect(ruta); dst=sqlite3.connect(destino)
            try: src.backup(dst)
            finally: src.close(); dst.close()
            datos=Path(destino).read_bytes()
            if not datos.startswith(b'SQLite format 3'): raise RuntimeError('Respaldo SQLite inválido')
            return datos,'sqlite3','application/octet-stream'
        finally: os.unlink(destino)
    ejecutable=shutil.which('pg_dump')
    if not ejecutable: raise RuntimeError('pg_dump no está instalado en el servidor. No se permitirá limpiar sin respaldo.')
    # pg_dump acepta URLs postgresql+psycopg tras normalización del dialecto.
    url=url.replace('postgresql+psycopg://','postgresql://',1).replace('postgresql+psycopg2://','postgresql://',1)
    resultado=subprocess.run([ejecutable,'--format=custom','--no-owner','--no-acl','--dbname',url],capture_output=True,timeout=180)
    if resultado.returncode or not resultado.stdout.startswith(b'PGDMP'):
        raise RuntimeError('Falló la creación o validación del respaldo PostgreSQL.')
    return resultado.stdout,'dump','application/octet-stream'

# V7.34: migración explícita de códigos históricos, sin alterar PK ni reparaciones.
PREFIJOS_COMPONENTES = {
    'CUERPO':'CU', 'EJE':'EJ', 'IMPULSOR':'IM', 'TAPA':'TA',
    'SOPORTE':'SO', 'FRAME':'FR', 'VOLUTA':'VO', 'TUERCA':'TU',
    'CAMISA':'CA', 'BASE':'BA', 'ACOPLAMIENTO':'AC', 'RODAMIENTO':'RO',
    'SELLO':'SE', 'MOTOR':'MO', 'DESCARGA':'DE', 'VALVULA':'VA'
}

def prefijo_componente_v734(nombre):
    normal=unicodedata.normalize('NFKD',nombre or '').encode('ascii','ignore').decode().upper()
    palabras=re.findall(r'[A-Z]+',normal)
    for palabra in palabras:
        if palabra in PREFIJOS_COMPONENTES: return PREFIJOS_COMPONENTES[palabra]
    return prefijo_codigo(palabras[0] if palabras else 'XX',2)

def propuesta_codigos_v734():
    """Determinista: orden por ID; contador separado por componente+marca."""
    filas=[]; contadores={}
    for c,m in db.session.query(ComponenteModelo,ModeloEquipo).join(ModeloEquipo,ComponenteModelo.modelo_id==ModeloEquipo.id).order_by(ComponenteModelo.id).all():
        base=f'{prefijo_componente_v734(c.nombre)}-{prefijo_codigo(m.marca,3)}'
        contadores[base]=contadores.get(base,0)+1
        codigo=f'{base}-{contadores[base]:04d}'
        filas.append({'id':c.id,'modelo_id':m.id,'modelo':m.modelo,'marca':m.marca,'nombre':c.nombre,'anterior':c.codigo or '', 'nuevo':codigo, 'reparaciones':db.session.query(ReparacionMaestra).filter_by(componente_modelo_id=c.id).count()})
    return filas

def huella_codigos_v734(filas):
    # Incluye códigos actuales y datos que determinan los nuevos códigos.
    contenido=[[f[k] for k in ('id','modelo_id','marca','modelo','nombre','anterior','nuevo')] for f in filas]
    return hashlib.sha256(json.dumps(contenido,ensure_ascii=False).encode('utf-8')).hexdigest()

@app.route('/administracion/codigos-componentes')
@login_required
def administrar_codigos_componentes():
    if not usuario_actual().es_admin: abort(403)
    filas=propuesta_codigos_v734()
    codigos=[x['nuevo'] for x in filas]
    duplicados=len(codigos)!=len(set(codigos))
    huella=huella_codigos_v734(filas)
    listo=session.get('codigos_v734_huella')==huella and bool(session.get('codigos_v734_token'))
    return render_template('administrar_codigos_componentes.html',filas=filas,duplicados=duplicados,listo=listo,habilitado=os.environ.get('HABILITAR_MIGRACION_CODIGOS')=='SI')

@app.route('/administracion/codigos-componentes/respaldo',methods=['POST'])
@login_required
def respaldo_codigos_componentes():
    if not usuario_actual().es_admin: abort(403)
    filas=propuesta_codigos_v734()
    if len({f['nuevo'] for f in filas})!=len(filas): abort(409)
    try: datos,ext,mime=respaldo_completo()
    except Exception as exc:
        flash(f'No fue posible generar respaldo: {exc}')
        return redirect(url_for('administrar_codigos_componentes'))
    session['codigos_v734_huella']=huella_codigos_v734(filas)
    session['codigos_v734_token']=secrets.token_urlsafe(24)
    audit('Respaldo previo a migración de códigos',f'{len(datos)} bytes; {len(filas)} componentes')
    return send_file(BytesIO(datos),mimetype=mime,as_attachment=True,download_name=f'INGEPRO_pre_codigos_{datetime.utcnow():%Y%m%d_%H%M%S}.{ext}')

@app.route('/administracion/codigos-componentes/aplicar',methods=['POST'])
@login_required
def aplicar_codigos_componentes():
    if not usuario_actual().es_admin: abort(403)
    if os.environ.get('HABILITAR_MIGRACION_CODIGOS')!='SI':
        flash('Migración bloqueada: habilite HABILITAR_MIGRACION_CODIGOS=SI únicamente después de verificar el respaldo.')
        return redirect(url_for('administrar_codigos_componentes'))
    if request.form.get('confirmacion')!='ACTUALIZAR CODIGOS':
        flash('Confirmación incorrecta.'); return redirect(url_for('administrar_codigos_componentes'))
    if not session.get('codigos_v734_token'):
        flash('Primero descargue un respaldo.'); return redirect(url_for('administrar_codigos_componentes'))
    try:
        # Serializa la migración frente a otra migración simultánea en PostgreSQL.
        if db.engine.dialect.name=='postgresql':
            db.session.execute(db.text('SELECT pg_advisory_xact_lock(7342026)'))
        filas=propuesta_codigos_v734()
        if huella_codigos_v734(filas)!=session.get('codigos_v734_huella'):
            db.session.rollback(); session.pop('codigos_v734_token',None)
            flash('Los componentes cambiaron desde el respaldo. Descargue uno nuevo.')
            return redirect(url_for('administrar_codigos_componentes'))
        if len(set(x['nuevo'] for x in filas))!=len(filas):
            db.session.rollback(); abort(409,'Hay códigos duplicados en la propuesta.')
        for fila in filas:
            c=db.session.get(ComponenteModelo,fila['id'])
            c.codigo=fila['nuevo']
            # OTComponente conserva PK y FK; se actualiza el código visible de las OT vinculadas.
            db.session.query(OTComponente).filter_by(componente_modelo_id=c.id).update({'codigo':fila['nuevo']},synchronize_session=False)
        db.session.add(Auditoria(usuario=usuario_actual().nombre,accion='Migración códigos componentes V7.34',detalle=f'{len(filas)} componentes; huella {huella_codigos_v734(filas)}'))
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    finally:
        session.pop('codigos_v734_token',None)
        session.pop('codigos_v734_huella',None)
    flash(f'Códigos actualizados: {len(filas)} componentes. Reparaciones y relaciones de OT conservadas.')
    return redirect(url_for('administrar_codigos_componentes'))

@app.route('/administracion/limpieza-ot')
@login_required
def limpieza_ot():
    if not usuario_actual().es_admin: abort(403)
    return render_template('limpieza_ot.html',conteos=conteos_limpieza(),respaldo_listo=bool(session.get('respaldo_ot_token')))

@app.route('/administracion/limpieza-ot/respaldo',methods=['POST'])
@login_required
def descargar_respaldo_ot():
    if not usuario_actual().es_admin: abort(403)
    try: datos,ext,mime=respaldo_completo()
    except Exception as e:
        flash(f'No se pudo generar el respaldo: {e}'); return redirect(url_for('limpieza_ot'))
    session['respaldo_ot_token']=secrets.token_urlsafe(24)
    session['respaldo_ot_conteos']=[list(x) for x in conteos_limpieza()]
    audit('Respaldo previo a limpieza generado',f'{len(datos)} bytes')
    return send_file(BytesIO(datos),mimetype=mime,as_attachment=True,download_name=f'INGEPRO_respaldo_{datetime.utcnow():%Y%m%d_%H%M%S}.{ext}')

@app.route('/administracion/limpieza-ot/ejecutar',methods=['POST'])
@login_required
def ejecutar_limpieza_ot():
    if not usuario_actual().es_admin: abort(403)
    if os.environ.get('HABILITAR_LIMPIEZA_OT')!='SI':
        flash('La limpieza está bloqueada hasta validar el respaldo y habilitarla expresamente en Render.')
        return redirect(url_for('limpieza_ot'))
    if request.form.get('confirmacion')!='ELIMINAR OT' or not session.get('respaldo_ot_token'):
        flash('Primero genere el respaldo y escriba ELIMINAR OT.'); return redirect(url_for('limpieza_ot'))
    if session.get('respaldo_ot_conteos')!=[list(x) for x in conteos_limpieza()]:
        session.pop('respaldo_ot_token',None)
        flash('Los datos cambiaron desde el respaldo. Genere uno nuevo.'); return redirect(url_for('limpieza_ot'))
    # Protección contra limpieza de una base distinta de la respaldada.
    try:
        for cls in TABLAS_LIMPIEZA: db.session.query(cls).delete(synchronize_session=False)
        db.session.add(Auditoria(usuario=usuario_actual().nombre,accion='Limpieza total de OT de prueba',detalle='Datos transaccionales eliminados, maestros conservados.'))
        db.session.commit()
    except Exception:
        db.session.rollback(); raise
    session.pop('respaldo_ot_token',None); session.pop('respaldo_ot_conteos',None)
    flash('Limpieza completada. Usuarios, clientes, equipos y componentes conservados.')
    return redirect(url_for('limpieza_ot'))



# INGEPRO RP V7.45 - Ordenes de trabajo internas
class OTInterna(db.Model):
    id=db.Column(db.Integer,primary_key=True)
    ot_id=db.Column(db.Integer,db.ForeignKey('ot.id'),nullable=False,index=True)
    numero=db.Column(db.String(40),unique=True,nullable=False)
    revision=db.Column(db.Integer,default=1,nullable=False)
    creado=db.Column(db.DateTime,default=datetime.utcnow)
    actualizado=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow)
    creado_por=db.Column(db.String(120))
    estado=db.Column(db.String(30),default='Borrador')
    contenido_json=db.Column(db.Text,nullable=False,default='[]')
    observaciones=db.Column(db.Text)
    ot=db.relationship('OT')


def oti_origen(ot):
    """Solo componentes recibidos y trabajos seleccionados o programados."""
    resultado=[]
    for comp in ot.componentes_detalle:
        if (comp.condicion_ingreso or '').strip().casefold()!='recibido': continue
        selecciones=OTReparacionSeleccionada.query.filter_by(ot_id=ot.id,ot_componente_id=comp.id).all()
        asignaciones=PlanActividad.query.filter_by(ot_id=ot.id,ot_componente_id=comp.id).order_by(PlanActividad.fecha_inicio,PlanActividad.id).all()
        trabajos={}
        for s in selecciones:
            rep=db.session.get(ReparacionMaestra,s.reparacion_maestra_id)
            if rep: trabajos['r'+str(rep.id)]={'nombre':rep.nombre,'instrucciones':rep.descripcion or '', 'horas':float(rep.hh_estimadas or 0),'responsable':'','estado':'Pendiente','origen':rep.id}
        for a in asignaciones:
            key='r'+str(a.reparacion_maestra_id) if a.reparacion_maestra_id else 'a'+str(a.id)
            if key not in trabajos:
                trabajos[key]={'nombre':a.reparacion.nombre if a.reparacion else a.tarea,'instrucciones':'','horas':0,'responsable':'','estado':'Pendiente','origen':a.reparacion_maestra_id}
            t=trabajos[key]
            if a.responsable: t['responsable']=a.responsable.nombre
            if not t['horas']: t['horas']=float((a.fecha_fin-a.fecha_inicio).days+1)*float(a.horas_dia or 0)
            if a.observacion and not t['instrucciones']: t['instrucciones']=a.observacion
        if trabajos: resultado.append({'componente_id':comp.id,'componente':comp.nombre,'etapas':list(trabajos.values())})
    return resultado

@app.route('/ot-internas')
@permiso_required('planificacion')
def oti_lista():
    buscar=(request.args.get('buscar') or '').strip()[:120]
    estado=(request.args.get('estado') or '').strip()
    consulta=OTInterna.query.join(OT,OTInterna.ot_id==OT.id)
    if estado in ('Borrador','Emitida'):
        consulta=consulta.filter(OTInterna.estado==estado)
    else:
        estado=''
    if buscar:
        from sqlalchemy import or_
        termino=f'%{buscar}%'
        consulta=consulta.filter(or_(OTInterna.numero.ilike(termino),OT.numero.ilike(termino),OT.cliente.ilike(termino),OT.equipo.ilike(termino)))
    items=consulta.order_by(OTInterna.id.desc()).limit(200).all()
    return render_template('ot_internas.html',items=items,buscar=buscar,estado=estado)

@app.route('/ot-internas/nueva',methods=['GET'])
@permiso_required('planificacion')
def oti_seleccionar():
    buscar=(request.args.get('buscar') or '').strip()[:120]
    consulta=OT.query
    if buscar:
        from sqlalchemy import or_
        termino=f'%{buscar}%'
        consulta=consulta.filter(or_(OT.numero.ilike(termino),OT.cliente.ilike(termino),OT.equipo.ilike(termino)))
    ots=consulta.order_by(OT.fecha.desc(),OT.id.desc()).limit(300).all()
    # Se cuentan solo los componentes recibidos; la OT se crea usando oti_origen.
    opciones=[]
    for ot in ots:
        recibidos=[c for c in ot.componentes_detalle if (c.condicion_ingreso or '').strip().casefold()=='recibido']
        opciones.append({'ot':ot,'componentes':len(recibidos)})
    return render_template('ot_interna_nueva.html',opciones=opciones,buscar=buscar)

@app.route('/ot-internas/nueva/<int:ot_id>',methods=['POST'])
@permiso_required('planificacion')
def oti_nueva(ot_id):
    ot=db.session.get(OT,ot_id) or abort(404)
    contenido=oti_origen(ot)
    if not contenido:
        flash('La OT no tiene trabajos seleccionados o planificados en componentes recibidos.')
        return redirect(url_for('oti_lista'))
    nuevo=OTInterna(ot_id=ot.id,numero='TEMP-'+secrets.token_hex(8),creado_por=usuario_actual().nombre,contenido_json=json.dumps(contenido,ensure_ascii=False))
    db.session.add(nuevo);db.session.flush()
    nuevo.numero='OTI-'+str(nuevo.id).zfill(6)
    db.session.commit()
    audit('OT interna creada',nuevo.numero+' / '+ot.numero)
    return redirect(url_for('oti_editar',id=nuevo.id))

@app.route('/ot-internas/<int:id>',methods=['GET','POST'])
@permiso_required('planificacion')
def oti_editar(id):
    doc=db.session.get(OTInterna,id) or abort(404)
    contenido=json.loads(doc.contenido_json or '[]')
    if request.method=='POST':
        for ci,c in enumerate(contenido):
            for ei,e in enumerate(c['etapas']):
                prefix=f'e_{ci}_{ei}_'
                e['nombre']=(request.form.get(prefix+'nombre') or e['nombre']).strip()[:180]
                e['instrucciones']=(request.form.get(prefix+'instrucciones') or '').strip()[:3000]
                e['responsable']=(request.form.get(prefix+'responsable') or '').strip()[:120]
                e['estado']=request.form.get(prefix+'estado','Pendiente') if request.form.get(prefix+'estado') in ('Pendiente','En ejecución','Terminada') else 'Pendiente'
                try:e['horas']=max(0,float(request.form.get(prefix+'horas') or 0))
                except ValueError:e['horas']=0
        doc.contenido_json=json.dumps(contenido,ensure_ascii=False)
        doc.observaciones=(request.form.get('observaciones') or '')[:3000]
        doc.estado='Emitida' if request.form.get('accion')=='emitir' else 'Borrador'
        db.session.commit();audit('OT interna guardada',doc.numero)
        flash('OT interna guardada correctamente.')
        return redirect(url_for('oti_editar',id=id))
    trabajadores=Usuario.query.filter_by(activo=True,disponible_planificacion=True).order_by(Usuario.nombre).all()
    return render_template('ot_interna_editar.html',doc=doc,contenido=contenido,trabajadores=trabajadores)

@app.route('/ot-internas/<int:id>/pdf')
@permiso_required('planificacion')
def oti_pdf(id):
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import KeepTogether
    from xml.sax.saxutils import escape
    doc=db.session.get(OTInterna,id) or abort(404)
    contenido=json.loads(doc.contenido_json or '[]')
    buf=BytesIO()
    pdf=SimpleDocTemplate(buf,pagesize=A4,rightMargin=35,leftMargin=35,topMargin=103,bottomMargin=55,title=doc.numero)
    w=A4[0]-70
    normal=ParagraphStyle('oti-normal',fontName='Helvetica',fontSize=8.2,leading=11,spaceAfter=0)
    small=ParagraphStyle('oti-small',parent=normal,fontSize=7.2,leading=10,textColor=colors.HexColor('#555555'))
    bold=ParagraphStyle('oti-bold',parent=normal,fontName='Helvetica-Bold')
    story=[]
    datos=[['OT PRINCIPAL',doc.ot.numero,'CLIENTE',doc.ot.cliente],['EQUIPO',doc.ot.equipo,'ESTADO',doc.estado]]
    meta=Table([[Paragraph(escape(str(v)),bold if j%2==0 else normal) for j,v in enumerate(row)] for row in datos],colWidths=[w*.18,w*.32,w*.18,w*.32])
    meta.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F5F5F5')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#CCCCCC')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story.extend([meta,Spacer(1,13)])
    total=0
    for c in contenido:
        cab=Table([[Paragraph(escape(c['componente'].upper()),bold)]],colWidths=[w])
        cab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#E1E1E1')),('BOX',(0,0),(-1,-1),.4,colors.HexColor('#BBBBBB')),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
        rows=[[Paragraph('N.º',bold),Paragraph('ETAPA / INSTRUCCIONES',bold),Paragraph('RESPONSABLE',bold),Paragraph('HRS.',bold)]]
        for j,e in enumerate(c['etapas'],1):
            h=float(e.get('horas') or 0);total+=h
            desc=escape(e.get('instrucciones') or '')
            cell=Paragraph('<b>'+escape(e.get('nombre') or '')+'</b>'+(('<br/>'+desc) if desc else ''),normal)
            rows.append([str(j),cell,Paragraph(escape(e.get('responsable') or 'SIN ASIGNAR'),small),f'{h:g}'])
        table=Table(rows,colWidths=[w*.06,w*.61,w*.23,w*.10],repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#F1F1F1')),('LINEBELOW',(0,0),(-1,-1),.35,colors.HexColor('#D5D5D5')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        story.extend([KeepTogether([cab,table]),Spacer(1,10)])
    summary=Table([[Paragraph('TOTAL HORAS ESTIMADAS',bold),Paragraph(f'{total:g} h',bold)]],colWidths=[w*.8,w*.2])
    summary.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#EDEDED')),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]))
    # Leyenda de responsables presentes en esta OT; PDF en escala de grises.
    nombres={e.get('responsable','').strip() for c in contenido for e in c.get('etapas',[]) if e.get('responsable','').strip()}
    sin_planificar=any(not e.get('responsable','').strip() for c in contenido for e in c.get('etapas',[]))
    personal=Usuario.query.filter(Usuario.nombre.in_(nombres)).all() if nombres else []
    mapa={u.nombre:u for u in personal}
    leyenda=[]
    for nombre in sorted(nombres):
        u=mapa.get(nombre)
        leyenda.append(f'{(u.sigla or "—") if u else "—"} = {nombre}')
    if sin_planificar: leyenda.append('S/P = SIN PLANIFICAR (sin trabajador asignado)')
    if leyenda:
        story.extend([Spacer(1,10),Paragraph('LEYENDA DE RESPONSABLES Y PLANIFICACIÓN',bold),Spacer(1,5)])
        for item in leyenda: story.append(Paragraph(escape(item),small))
        story.append(Paragraph('Los colores de trabajadores se consultan en la OT digital y Planificación.',small))
    story.extend([summary,Spacer(1,14),Paragraph('OBSERVACIONES GENERALES',bold),Spacer(1,5),Paragraph(escape(doc.observaciones or '____________________________________________________________'),normal),Spacer(1,54)])
    firmas=Table([['____________________________','____________________________'],['JEFE DE ÁREA','RECEPCIÓN DE TRABAJOS'],['Nombre, firma y fecha','Nombre, firma y fecha']],colWidths=[w/2,w/2])
    firmas.setStyle(TableStyle([('ALIGN',(0,0),(-1,-1),'CENTER'),('FONTNAME',(0,1),(-1,1),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),8),('TEXTCOLOR',(0,2),(-1,2),colors.grey),('BOTTOMPADDING',(0,0),(-1,0),8)]))
    story.append(firmas)
    def pagina(canv,pdfdoc):
        canv.saveState();pw,ph=A4
        canv.setFillColor(colors.HexColor('#D4D4D4'));canv.rect(35,ph-91,pw-70,59,stroke=0,fill=1)
        logo=Path(app.root_path)/'static/img/logo-ingepro-small.png'
        if logo.exists():
            try:
                from PIL import Image as Img,ImageOps as Ops
                src=Img.open(logo).convert('RGBA'); alpha=src.getchannel('A');g=Ops.grayscale(src.convert('RGB'));mono=Img.merge('RGBA',(g,g,g,alpha))
                canv.drawImage(ImageReader(mono),44,ph-83,width=140,height=43,mask='auto',preserveAspectRatio=True)
            except Exception: pass
        canv.setStrokeColor(colors.HexColor('#A0A0A0'));canv.line(197,ph-83,197,ph-41)
        canv.setFillColor(colors.HexColor('#222222'));canv.setFont('Helvetica-Bold',12);canv.drawString(211,ph-54,'ORDEN DE TRABAJO INTERNA')
        canv.setFont('Helvetica',9);canv.drawString(211,ph-69,'DIVISIÓN MAESTRANZA  |  '+doc.numero)
        canv.setStrokeColor(colors.HexColor('#CCCCCC'));canv.line(35,42,pw-35,42)
        canv.setFont('Helvetica',7);canv.drawString(35,30,'INGEPRO · DIVISIÓN MAESTRANZA');canv.drawRightString(pw-35,30,f'Página {pdfdoc.page}')
        canv.restoreState()
    pdf.build(story,onFirstPage=pagina,onLaterPages=pagina)
    buf.seek(0)
    return send_file(buf,mimetype='application/pdf',as_attachment=False,download_name=doc.numero+'.pdf')


@app.route('/produccion')
@permiso_required('produccion')
def produccion():
    buscar=(request.args.get('buscar') or '').strip()
    filas=OTReparacionSeleccionada.query.order_by(OTReparacionSeleccionada.ot_id.desc(),OTReparacionSeleccionada.ot_componente_id).all()
    if buscar:
        ots=OT.query.filter((OT.numero.ilike('%'+buscar+'%')) | (OT.equipo.ilike('%'+buscar+'%')) | (OT.cliente.ilike('%'+buscar+'%'))).with_entities(OT.id).all()
        ids=[r[0] for r in ots]
        filas=[r for r in filas if r.ot_id in ids]
    ids_ot={r.ot_id for r in filas}; ids_comp={r.ot_componente_id for r in filas}; ids_rep={r.reparacion_maestra_id for r in filas}
    ots={o.id:o for o in OT.query.filter(OT.id.in_(ids_ot)).all()} if ids_ot else {}
    comps={c.id:c for c in OTComponente.query.filter(OTComponente.id.in_(ids_comp)).all()} if ids_comp else {}
    reps={r.id:r for r in ReparacionMaestra.query.filter(ReparacionMaestra.id.in_(ids_rep)).all()} if ids_rep else {}
    etapas=ProduccionEtapa.query.filter(ProduccionEtapa.ot_componente_id.in_(ids_comp)).all() if ids_comp else []
    estados={(e.ot_componente_id,e.reparacion_maestra_id):e for e in etapas}
    actividades=PlanActividad.query.filter(PlanActividad.ot_componente_id.in_(ids_comp)).all() if ids_comp else []
    plan={}
    for a in actividades:
        plan.setdefault((a.ot_componente_id,a.reparacion_maestra_id),[]).append(a)
    grupos={}
    for f in filas:
        ot=ots.get(f.ot_id); comp=comps.get(f.ot_componente_id); rep=reps.get(f.reparacion_maestra_id)
        if not ot or not comp or not rep or comp.condicion_ingreso!='Recibido': continue
        key=(ot.id,comp.id)
        if key not in grupos: grupos[key]={'ot':ot,'componente':comp,'etapas':[]}
        registro=estados.get((comp.id,rep.id))
        asignaciones=plan.get((comp.id,rep.id),[])
        estado=registro.estado if registro else ('En ejecución' if asignaciones and any(a.estado in ('En ejecución','En proceso') for a in asignaciones) else 'Pendiente')
        grupos[key]['etapas'].append({'rep':rep,'estado':estado,'registro':registro,'asignaciones':asignaciones})
    ordenes={}
    for (ot_id,comp_id),g in grupos.items():
        g['etapas'].sort(key=lambda e:(e['rep'].secuencia or 0,e['rep'].id))
        if ot_id not in ordenes:
            ordenes[ot_id]={'ot':g['ot'],'componentes':[],'total':0,'terminadas':0}
        ordenes[ot_id]['componentes'].append(g)
        ordenes[ot_id]['total']+=len(g['etapas'])
        ordenes[ot_id]['terminadas']+=sum(e['estado']=='Terminada' for e in g['etapas'])
    for orden in ordenes.values():
        orden['componentes'].sort(key=lambda g:g['componente'].id)
    if 'produccion_token' not in session: session['produccion_token']=secrets.token_urlsafe(24)
    return render_template('produccion.html',ordenes=list(ordenes.values()),buscar=buscar,gerencia=usuario_actual().es_admin,token=session['produccion_token'])

@app.route('/produccion/terminar/<int:ot_id>/<int:comp_id>/<int:rep_id>',methods=['POST'])
@permiso_required('produccion')
def produccion_terminar(ot_id,comp_id,rep_id):
    if not secrets.compare_digest(request.form.get('token',''),session.get('produccion_token','')): abort(400)
    seleccionado=OTReparacionSeleccionada.query.filter_by(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id).first_or_404()
    comp=OTComponente.query.filter_by(id=comp_id,ot_id=ot_id,condicion_ingreso='Recibido').first_or_404()
    rep=ReparacionMaestra.query.filter_by(id=rep_id,componente_modelo_id=comp.componente_modelo_id).first_or_404()
    asignaciones=PlanActividad.query.filter_by(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id).all()
    if not asignaciones:
        flash('Esta etapa no tiene planificación. No se puede cerrar.')
        return redirect(url_for('produccion'))
    etapa=ProduccionEtapa.query.filter_by(ot_componente_id=comp_id,reparacion_maestra_id=rep_id).first()
    if etapa and etapa.estado=='Terminada':
        flash('Esta etapa ya fue terminada.')
        return redirect(url_for('produccion'))
    anteriores=OTReparacionSeleccionada.query.filter_by(ot_id=ot_id,ot_componente_id=comp_id).all()
    prev_ids=[r.reparacion_maestra_id for r in anteriores]
    prev_reps=ReparacionMaestra.query.filter(ReparacionMaestra.id.in_(prev_ids)).order_by(ReparacionMaestra.secuencia,ReparacionMaestra.id).all()
    for anterior in prev_reps:
        if anterior.id==rep_id: break
        prev_estado=ProduccionEtapa.query.filter_by(ot_componente_id=comp_id,reparacion_maestra_id=anterior.id).first()
        if not prev_estado or prev_estado.estado!='Terminada':
            flash('Primero debe finalizar la etapa anterior: '+anterior.nombre)
            return redirect(url_for('produccion'))
    if not etapa:
        etapa=ProduccionEtapa(ot_id=ot_id,ot_componente_id=comp_id,reparacion_maestra_id=rep_id)
        db.session.add(etapa)
    etapa.estado='Terminada'; etapa.observacion=(request.form.get('observacion') or '').strip()[:3000]
    etapa.terminado_por_id=usuario_actual().id; etapa.terminado_fecha=datetime.utcnow()
    db.session.commit()
    audit('Etapa de reparación terminada',f'OT {ot_id}; componente {comp_id}; reparación {rep.nombre}; jefatura {usuario_actual().nombre}')
    flash('Etapa terminada y registrada. La siguiente etapa queda disponible según su planificación.')
    return redirect(url_for('produccion',buscar=request.form.get('buscar','')))

with app.app_context(): db.create_all(); asegurar_columnas(); seed_maestros()
if __name__=='__main__': app.run(host='0.0.0.0',port=int(os.environ.get('PORT',5000)),debug=True)
